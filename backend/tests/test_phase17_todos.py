"""Phase 17: daily to-do lists. Supervisors, foremen, managers and the office set them; nobody clocks off until every task
is ticked off or has a reason it couldn't be finished."""
from datetime import timedelta

from app import todos as td
from tests.test_phase16_tools_safety import _notes, _user


def test_todos_block_clock_off_until_ticked_or_explained(client):
    fore, fh = _user(client, "todo-fore@test.local", name="Fred Foreman", role="field", trade="foreman")
    ben, bh = _user(client, "todo-ben@test.local", name="Ben Plumber", role="field", trade="plumber")
    amy, ah = _user(client, "todo-amy@test.local", name="Amy Apprentice", role="field", trade="apprentice")
    # The foreman sets the day's list for two of the crew; they're told
    r = client.post("/api/field/todos", headers=fh, json={"user_ids": [ben, amy], "texts": ["Pressure test level 2", "Clean up the slab", " "]})
    assert r.status_code == 200 and len(r.json()) == 4
    assert _notes(ben, "Fred Foreman set 2 to-dos for you today")
    # A plumber can keep his own list, but can't set one for anyone else
    assert client.post("/api/field/todos", headers=bh, json={"texts": ["Order 100 mm bends"]}).status_code == 200
    assert client.post("/api/field/todos", headers=bh, json={"user_ids": [amy], "texts": ["Sweep up"]}).status_code == 403
    mine = client.get("/api/field/todos", headers=bh).json()
    assert [t["text"] for t in mine["tasks"]] == ["Pressure test level 2", "Clean up the slab", "Order 100 mm bends"] and mine["open"] == 3
    assert mine["leads_team"] is False and client.get("/api/field/todos", headers=fh).json()["leads_team"] is True

    # Ben clocks on, then can't clock off with tasks still open
    client.post("/api/field/clock-in", headers=bh, json={})
    r = client.post("/api/field/clock-out", headers=bh, json={})
    assert r.status_code == 409 and "to-dos" in r.json()["detail"]["message"] and "Pressure test level 2" in r.json()["detail"]["message"]
    test, slab, order = mine["tasks"]
    client.patch(f"/api/field/todos/{test['id']}", headers=bh, json={"status": "done"})
    # "Couldn't finish" needs a reason, which goes back to the foreman
    assert client.patch(f"/api/field/todos/{slab['id']}", headers=bh, json={"status": "not_done", "comment": ""}).status_code == 422
    client.patch(f"/api/field/todos/{slab['id']}", headers=bh, json={"status": "not_done", "comment": "Concrete truck late, slab still wet"})
    assert _notes(fore, "slab still wet")
    # He can't delete a task the foreman set, only his own
    assert client.delete(f"/api/field/todos/{test['id']}", headers=bh).status_code == 403
    assert client.delete(f"/api/field/todos/{order['id']}", headers=bh).status_code == 204
    assert client.post("/api/field/clock-out", headers=bh, json={}).status_code == 200

    # The crew clock-off skips anyone whose list isn't finished
    client.post("/api/field/clock-in", headers=ah, json={})
    r = client.post("/api/field/crew", headers=fh, json={"user_ids": [amy], "out": True}).json()
    assert r["done"] == [] and r["skipped"] == ["Amy Apprentice"]

    # The foreman sees the whole crew's day, then carries the unfinished work to tomorrow
    team = {p["name"]: p for p in client.get("/api/field/todos/team", headers=fh).json()["people"]}
    assert team["Ben Plumber"]["done"] == 1 and team["Ben Plumber"]["not_done"] == 1 and team["Amy Apprentice"]["open"] == 2
    assert team["Ben Plumber"]["tasks"][1]["comment"] == "Concrete truck late, slab still wet"
    today = td.today()
    r = client.post("/api/field/todos/carry", headers=fh, json={"from_day": str(today), "to_day": str(today + timedelta(days=1)), "user_ids": [ben, amy]})
    assert r.json()["carried"] == 3  # Ben's slab, and both of Amy's untouched ones
    assert client.post("/api/field/todos/carry", headers=fh, json={"from_day": str(today), "to_day": str(today + timedelta(days=1)),
                                                                   "user_ids": [ben, amy]}).json()["carried"] == 0  # not twice
    tomorrow = client.get(f"/api/field/todos?day={today + timedelta(days=1)}", headers=bh).json()["tasks"]
    assert [t["text"] for t in tomorrow] == ["Clean up the slab"] and tomorrow[0]["set_by"] == "Fred Foreman"


def test_managers_and_the_office_set_lists_for_their_people(client):
    boss, bh = _user(client, "todo-boss@test.local", name="Pat Plumber", role="field", trade="plumber")
    kid, kh = _user(client, "todo-kid@test.local", name="Kid Apprentice", role="field", trade="apprentice")
    other, _ = _user(client, "todo-other@test.local", name="Other Labourer", role="field", trade="labourer")
    client.put(f"/api/people/{kid}/profile", json={"manager_id": boss})
    # A plumber with an apprentice reporting to him can set the apprentice's list, and only his
    assert client.get("/api/field/todos", headers=bh).json()["leads_team"] is True
    assert client.post("/api/field/todos", headers=bh, json={"user_ids": [kid], "texts": ["Cut and deburr the copper"]}).status_code == 200
    assert client.post("/api/field/todos", headers=bh, json={"user_ids": [other], "texts": ["x"]}).status_code == 403
    names = [p["name"] for p in client.get("/api/field/todos/team", headers=bh).json()["people"]]
    assert names == ["Pat Plumber", "Kid Apprentice"]
    # The office (an admin here) can set anyone's, for tomorrow too
    tomorrow = td.today() + timedelta(days=1)
    r = client.post("/api/field/todos", json={"user_ids": [other], "texts": ["Collect the skip permit"], "day": str(tomorrow)})
    assert r.status_code == 200 and _notes(other, "for you tomorrow")
    # Nothing on today's list: clocking off works as before
    client.post("/api/field/clock-in", headers=kh, json={})
    item = client.get("/api/field/todos", headers=kh).json()["tasks"][0]
    client.patch(f"/api/field/todos/{item['id']}", headers=kh, json={"status": "done", "comment": "Done, 40 lengths"})
    assert client.post("/api/field/clock-out", headers=kh, json={}).status_code == 200


def test_talk_a_list_in(client):
    from app import chat

    boss, bh = _user(client, "talk-boss@test.local", name="Gary Foreman", role="field", trade="foreman")
    dave, _ = _user(client, "talk-dave@test.local", name="Davo Pipes", role="field", trade="plumber")
    sam, _ = _user(client, "talk-sam@test.local", name="Sammy Carter", role="field", trade="apprentice")
    said = ("Davo, pressure test the level 2 stack and then clean up the slab. Sammy can you pack down the trench shield, "
            "and grab the gas detector off Lee. Everyone lock the site container. I'll ring the inspector")
    # Without the AI: names start each person's tasks, "and" / "then" split them
    chat.set_model(None)
    r = client.post("/api/field/todos/parse", headers=bh, json={"text": said}).json()
    assert r["used_ai"] is False
    got = {li["name"]: [t["text"] for t in li["tasks"]] for li in r["lists"]}
    assert got["Davo Pipes"] == ["Pressure test the level 2 stack", "Clean up the slab", "Lock the site container"]
    assert got["Sammy Carter"] == ["Pack down the trench shield", "Grab the gas detector off Lee", "Lock the site container"]
    assert got["Gary Foreman"] == ["Ring the inspector"]
    assert "tomorrow" not in said and r["day"] == str(td.today())
    # With the AI: whatever it hands back is checked against the crew and cleaned up
    class Fake:
        def summarise(self, prompt):
            assert "Davo Pipes" in prompt and "pressure test" in prompt
            return ('Here you go: {"day": null, "lists": [{"user_id": %d, "tasks": [{"text": "pressure test level 2", "job_id": 99999}]},'
                    ' {"user_id": 123456, "tasks": [{"text": "someone not on the crew"}]}], "unmatched": ["buy milk"]}' % dave)
    chat.set_model(Fake())
    try:
        r = client.post("/api/field/todos/parse", headers=bh, json={"text": said}).json()
    finally:
        chat.set_model(None)
    assert r["used_ai"] and r["lists"] == [{"user_id": dave, "name": "Davo Pipes", "is_me": False,
                                            "tasks": [{"text": "Pressure test level 2", "job_id": None, "job": ""}]}]
    assert r["unmatched"] == ["buy milk", "someone not on the crew"]
    # Checked and saved in one go; each person told once
    tomorrow = td.today() + timedelta(days=1)
    r = client.post("/api/field/todos/batch", headers=bh, json={"day": str(tomorrow), "lists": [
        {"user_id": dave, "tasks": [{"text": "Pressure test level 2"}, {"text": "Clean up the slab"}]},
        {"user_id": sam, "tasks": [{"text": "Pack down the trench shield"}]}]})
    assert r.json() == {"added": 3, "people": 2}
    assert _notes(dave, "set 2 to-dos for you tomorrow")
    plumber_h = _user(client, "talk-plumber@test.local", role="field", trade="plumber")[1]
    assert client.post("/api/field/todos/batch", headers=plumber_h, json={"day": str(tomorrow), "lists": [
        {"user_id": dave, "tasks": [{"text": "x"}]}]}).status_code == 403


def test_performance_report_by_week_and_month(client):
    from datetime import date as d_
    from app.db import SessionLocal
    from app.models import DailyTask

    boss, bh = _user(client, "perf-boss@test.local", name="Perf Boss", role="field", trade="supervisor")
    good, _ = _user(client, "perf-good@test.local", name="Good Worker", role="field", trade="plumber")
    poor, _ = _user(client, "perf-poor@test.local", name="Poor Worker", role="field", trade="labourer")
    t = td.today()
    with SessionLocal() as db:
        for back in range(1, 15):  # two weeks of history
            day = t - timedelta(days=back)
            for i in range(4):
                db.add(DailyTask(user_id=good, day=day, text=f"g{i}", set_by=boss, status="done" if i < 4 or back % 7 else "not_done"))
                st = "done" if i < 2 else ("not_done" if i == 2 else "open")
                db.add(DailyTask(user_id=poor, day=day, text=f"p{i}", set_by=boss, status=st, comment="Ran out of time" if st == "not_done" else ""))
        db.add(DailyTask(user_id=poor, day=t, text="today, still going", set_by=boss, status="open"))  # not late yet
        db.commit()
    start = t - timedelta(days=14)
    rep = client.get(f"/api/field/todos/report?start={start}&end={t}&group=week").json()
    rows = {r["name"]: r for r in rep["people"] if r["name"] in ("Good Worker", "Poor Worker")}
    g, p = rows["Good Worker"], rows["Poor Worker"]
    assert (g["tasks"], g["done"], g["done_pct"], g["rating"]) == (56, 56, 100.0, "Performing")
    assert (p["tasks"], p["done"], p["not_done"], p["untouched"], p["done_pct"], p["actioned_pct"]) == (56, 28, 14, 14, 50.0, 75.0)
    assert p["rating"] == "Needs attention" and p["reasons"][0]["reason"] == "Ran out of time" and p["days_all_done"] == 0
    assert rep["people"][0]["rating"] == "Needs attention"  # worst first
    assert len(rep["periods"]) >= 2 and all(k in p["periods"] for k in rep["periods"])
    monthly = client.get(f"/api/field/todos/report?start={start}&end={t}&group=month").json()
    assert all(len(k) == 7 for k in monthly["periods"])
    # The supervisor sees the people he looks after; a plumber's position can't open the report at all
    assert {"Good Worker", "Poor Worker"} <= {r["name"] for r in client.get("/api/field/todos/report", headers=bh).json()["people"]}
    pid, plumber_h = _user(client, "perf-plumber@test.local", role="field", trade="plumber")
    pos = {p["name"]: p for p in client.get("/api/positions").json()}
    client.put(f"/api/users/{pid}/access", json={"position_id": pos["Plumber"]["id"], "features": {}})
    assert client.get("/api/field/todos/report", headers=plumber_h).status_code == 403
    x = client.get(f"/api/field/todos/report.xlsx?start={start}&end={t}&group=month")
    assert x.status_code == 200 and x.content[:2] == b"PK"
    import io
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(x.content))
    assert wb.sheetnames == ["Summary", "By month", "Reasons"]
    names = [c.value for c in wb["Summary"]["A"]]
    assert "Poor Worker" in names and any(row[3].value == "Ran out of time" for row in wb["Reasons"].iter_rows(min_row=2))
    del d_
