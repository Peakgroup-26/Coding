"""Phase 3 acceptance: an imported price book reproduces a known estimate within ±1%, and
unmapped items block final pricing. Also covers versions, margins, allowances, Simpro CSV."""
import csv
import io

import pytest
from openpyxl import Workbook, load_workbook

from .helpers import calibrate, count, new_job, run, settings, upload

ITEMS = [("Part Number", "Description", "UOM", "Trade Price"),
         ("PVC100", "PVC DWV DN100 pipe", "m", "12.50"),
         ("PANCON", "Pan connector", "ea", "22.00"),
         ("CEMENT", "Solvent cement 500 mL", "tin", "18.00"),
         ("BASINWASTE", "Basin waste kit", "ea", "$31.00")]
RATES = "Code,Description,Charge Out Rate\nPLUMB,Plumber,95.00\nAPPR,Apprentice,55\n"
# Simpro-style long format: prebuild code repeated once per component; PB-WC nests PB-SEAL.
PREBUILDS = ("Prebuild Code,Prebuild Name,Labour Time,Component Code,Component Qty\n"
             "PB-DWV100M,Lay DN100 DWV per m,0.25,PVC100,1\n"
             "PB-WC,Install WC,1.5,PANCON,1\n"
             "PB-WC,,,PB-SEAL,1\n"
             "PB-SEAL,Seal kit,0.1,CEMENT,0.05\n")


def _xlsx(rows) -> bytes:
    wb = Workbook()
    for r in rows:
        wb.active.append(list(r))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _import(client, vid, kind, name, data):
    r = client.post(f"/api/price-books/{vid}/import", data={"kind": kind}, files={"file": (name, data)})
    assert r.status_code == 200, r.text
    assert r.json()["errors"] == [], r.json()
    return r.json()


@pytest.fixture(scope="module")
def priced(client):
    services, stages = settings(client)
    book = client.post("/api/price-books", json={"label": "Peak prices Oct 2026"}).json()
    vid = book["id"]
    assert _import(client, vid, "items", "prices.xlsx", _xlsx(ITEMS))["created"] == 4
    _import(client, vid, "rates", "rates.csv", RATES.encode())
    res = _import(client, vid, "prebuilds", "prebuilds.csv", PREBUILDS.encode())
    assert res["created"] == 3 and res["columns"]["component_code"] == "Component Code"
    client.patch(f"/api/price-books/{vid}", json={"default_labour_rate_code": "PLUMB"})

    job = new_job(client, name="Golden job", quote_no="Q-G1")
    sheet = upload(client, job["id"], [{"sheet_no": "H-950", "title": "SANITARY DRAINAGE PLAN"}])[0]
    calibrate(client, sheet["id"])
    san = services["SAN"]["id"]
    run(client, sheet["id"], san, 100, "PVC DWV", [(100, 100), (200, 100), (200, 150)])  # 15 m
    run(client, sheet["id"], san, 100, "PVC DWV", [(150, 130), (150, 100)])  # 3 m
    count(client, sheet["id"], san, "WC", 150, 130)

    for m in ({"source_kind": "pipe", "service_code": "SAN", "material": "PVC DWV", "dn": 100, "target_kind": "prebuild", "target_code": "PB-DWV100M"},
              {"source_kind": "count", "item_type": "WC", "target_kind": "prebuild", "target_code": "PB-WC"}):
        assert client.post("/api/mappings", json=m).status_code == 200

    r = client.put(f"/api/jobs/{job['id']}/pricing-settings", json={
        "price_book_version_id": vid, "margin_basis": "markup", "default_margin_pct": 20,
        "stage_margins": {"Fit-off": 30}, "gst_pct": 10, "reference_total_ex_gst": 3340})
    assert r.status_code == 200, r.text
    jid = job["id"]
    client.post(f"/api/jobs/{jid}/allowances", json={"kind": "pc", "description": "WC suite supply (PC)", "amount": 700})
    client.post(f"/api/jobs/{jid}/allowances", json={"kind": "provisional", "description": "Authority fees", "amount": 1500, "apply_margin": False})
    client.post(f"/api/jobs/{jid}/allowances", json={"kind": "exclusion", "description": "Builder's works and core holes"})
    return {"job": job, "vid": vid, "sheet": sheet, "services": services, "stages": stages}


def test_known_estimate_reproduced(client, priced):
    p = client.get(f"/api/jobs/{priced['job']['id']}/pricing").json()
    assert p["unmapped"] == [] and p["missing_codes"] == []
    lines = {l["target_code"]: l for l in p["lines"]}
    pipe = lines["PB-DWV100M"]
    # 18 m x ($12.50 + 0.25 h x $95)
    assert pipe["qty"] == 18 and pipe["material_cost"] == 225.00 and pipe["labour_cost"] == 427.50 and pipe["labour_hours"] == 4.5
    assert pipe["sell"] == 783.00  # +20%
    wc = lines["PB-WC"]
    # $22 + nested seal kit (0.05 x $18 = $0.90; 0.1 h) and 1.5 h + 0.1 h at $95
    assert wc["material_cost"] == 22.90 and wc["labour_cost"] == 152.00 and wc["labour_hours"] == 1.6
    assert wc["margin_pct"] == 30 and wc["sell"] == 227.37  # Fit-off stage margin
    s = p["summary"]
    assert s["cost"] == pytest.approx(827.40 + 2200)
    assert s["sell_ex_gst"] == pytest.approx(783.00 + 227.37 + 840.00 + 1500.00)  # PC +20%, provisional sum no margin
    assert s["gst"] == pytest.approx(335.04) and s["sell_inc_gst"] == pytest.approx(3685.41)
    # Acceptance: within ±1% of the known past estimate
    assert abs(s["reference_variance_pct"]) <= 1.0
    assert any(a["kind"] == "exclusion" for a in p["allowances"])
    stage_sell = sum(r["sell"] for r in s["by_stage"])
    assert stage_sell == pytest.approx(s["sell_ex_gst"])


def test_margin_basis(client, priced):
    jid = priced["job"]["id"]
    client.put(f"/api/jobs/{jid}/pricing-settings", json={"price_book_version_id": priced["vid"], "margin_basis": "margin",
                                                          "default_margin_pct": 20, "stage_margins": {"Fit-off": 30}, "reference_total_ex_gst": 3340})
    p = client.get(f"/api/jobs/{jid}/pricing").json()
    pipe = next(l for l in p["lines"] if l["target_code"] == "PB-DWV100M")
    assert pipe["sell"] == 815.63  # 652.50 / 0.8, to the cent
    assert client.put(f"/api/jobs/{jid}/pricing-settings", json={"margin_basis": "margin", "default_margin_pct": 100}).status_code == 422
    client.put(f"/api/jobs/{jid}/pricing-settings", json={"price_book_version_id": priced["vid"], "margin_basis": "markup",
                                                          "default_margin_pct": 20, "stage_margins": {"Fit-off": 30}, "reference_total_ex_gst": 3340})


def test_versions_keep_old_rates(client, priced):
    vid, jid = priced["vid"], priced["job"]["id"]
    before = client.get(f"/api/jobs/{jid}/pricing").json()["summary"]["sell_ex_gst"]
    assert client.post(f"/api/price-books/{vid}/lock").status_code == 200
    # Locked books can't change
    assert client.put(f"/api/price-books/{vid}/items", json={"code": "PVC100", "cost": 99}).status_code == 409
    v2 = client.post("/api/price-books", json={"label": "Peak prices Nov 2026", "copy_from_id": vid}).json()
    assert client.put(f"/api/price-books/{v2['id']}/items", json={"code": "PVC100", "description": "PVC DWV DN100 pipe", "uom": "m", "cost": 14.0}).status_code == 200
    # The job still uses the old version
    assert client.get(f"/api/jobs/{jid}/pricing").json()["summary"]["sell_ex_gst"] == before
    books = {b["id"]: b for b in client.get("/api/price-books").json()}
    assert books[v2["id"]]["items"] == 4 and books[v2["id"]]["prebuilds"] == 3 and books[vid]["jobs"] == 1


def test_simpro_csv_and_template(client, priced):
    jid = priced["job"]["id"]
    r = client.get(f"/api/jobs/{jid}/simpro.csv")
    assert r.status_code == 200
    rows = list(csv.DictReader(io.StringIO(r.text)))
    assert rows[0].keys() >= {"Part Number", "Description", "Quantity", "Unit Sell Price", "Cost Centre"}
    pipe = next(x for x in rows if x["Part Number"] == "PB-DWV100M")
    assert float(pipe["Quantity"]) == 18 and pipe["Cost Centre"] == "Under slab" and pipe["Type"] == "Prebuild"
    assert any(x["Type"] == "PC item" for x in rows)
    # Admin changes the layout
    t = client.put("/api/settings/simpro-template", json={"columns": [{"header": "Code", "field": "code"}, {"header": "Qty", "field": "qty"},
                                                                      {"header": "Sell", "field": "sell_ex"}],
                                                          "group_by_stage": False, "include_allowances": False})
    assert t.status_code == 200
    rows = list(csv.DictReader(io.StringIO(client.get(f"/api/jobs/{jid}/simpro.csv").text)))
    assert list(rows[0].keys()) == ["Code", "Qty", "Sell"] and len(rows) == 2
    assert client.put("/api/settings/simpro-template", json={"columns": [{"header": "X", "field": "nope"}]}).status_code == 422


def test_unmapped_blocks_final_pricing(client, priced):
    jid, sid = priced["job"]["id"], priced["sheet"]["id"]
    basin = count(client, sid, priced["services"]["SAN"]["id"], "Basin", 400, 300)
    p = client.get(f"/api/jobs/{jid}/pricing").json()
    assert [u["item_type"] for u in p["unmapped"]] == ["Basin"]
    assert client.get(f"/api/jobs/{jid}/simpro.csv").status_code == 409

    flags = client.post(f"/api/jobs/{jid}/checks").json()
    assert any(f["check_type"] == "unmapped" and f["severity"] == "error" for f in flags)
    # Even with every flag accepted, sign-off refuses while something is unmapped
    for f in flags:
        client.patch(f"/api/flags/{f['id']}", json={"status": "accepted", "note": "ok"})
    r = client.post(f"/api/jobs/{jid}/finalise")
    assert r.status_code == 409 and "unmapped" in r.json()["detail"]

    # Map it to a catalogue item that isn't in the price book: still blocked as a missing code
    m = client.post("/api/mappings", json={"source_kind": "count", "item_type": "Basin", "target_kind": "item", "target_code": "NOPE"}).json()
    p = client.get(f"/api/jobs/{jid}/pricing").json()
    assert p["unmapped"] == [] and p["missing_codes"] == ["NOPE"]
    client.put(f"/api/mappings/{m['id']}", json={"source_kind": "count", "item_type": "Basin", "target_kind": "item", "target_code": "BASINWASTE"})

    for f in client.post(f"/api/jobs/{jid}/checks").json():
        if f["status"] == "open":
            client.patch(f"/api/flags/{f['id']}", json={"status": "accepted", "note": "ok"})
    r = client.post(f"/api/jobs/{jid}/finalise")
    assert r.status_code == 200, r.text
    assert client.post(f"/api/jobs/{jid}/allowances", json={"kind": "pc", "description": "late", "amount": 1}).status_code == 409
    client.post(f"/api/jobs/{jid}/reopen")
    client.delete(f"/api/counts/{basin['id']}")


def test_golden_report_and_export(client, priced):
    jid = priced["job"]["id"]
    ref = "service,dn,material,qty\nSAN,100,PVC DWV,18.5\nCW,,,10\n"
    assert client.post(f"/api/jobs/{jid}/reference", files={"file": ("ref.csv", ref.encode())}).json()["lines"] == 2
    rep = client.get(f"/api/jobs/{jid}/golden-report").json()
    svc = {s["service_code"]: s for s in rep["services"]}
    assert svc["SAN"]["reference_m"] == 18.5 and svc["SAN"]["takeoff_m"] == 18 and svc["SAN"]["variance_pct"] == pytest.approx(-2.7, abs=0.01)
    assert svc["CW"]["takeoff_m"] == 0
    assert abs(rep["price"]["variance_pct"]) <= 1

    wb = load_workbook(io.BytesIO(client.get(f"/api/jobs/{jid}/export.xlsx").content))
    assert "Pricing" in wb.sheetnames
    vals = [r for r in wb["Pricing"].iter_rows(values_only=True)]
    assert any(r and r[0] == "Total ex GST" for r in vals)


def test_import_errors_reported(client):
    book = client.post("/api/price-books", json={"label": "Bad import"}).json()
    r = client.post(f"/api/price-books/{book['id']}/import", data={"kind": "items"},
                    files={"file": ("x.csv", b"Code,Cost\nA,1.5\nB,abc\n")})
    assert r.json()["created"] == 1 and len(r.json()["errors"]) == 1
    r = client.post(f"/api/price-books/{book['id']}/import", data={"kind": "items"}, files={"file": ("x.csv", b"Foo,Bar\n1,2\n")})
    assert r.status_code == 422
    # Explicit column names
    r = client.post(f"/api/price-books/{book['id']}/import", data={"kind": "items", "columns": '{"code": "Foo", "cost": "Bar"}'},
                    files={"file": ("x.csv", b"Foo,Bar\nZ,2\n")})
    assert r.json()["created"] == 1
