"""Phase 10: material lists, labour hours, crew predictions, programs, cash flow and the dashboard."""
import io
from datetime import date

import pytest
from openpyxl import load_workbook

from app import ai
from app.reports import DEFAULTS, Calendar, crew_for, holidays, plumbers_for

from .test_phase7_autopilot import FakeReader, _plan_pdf
from .test_phase7_civil import FakeCivilReader, _civil_pdf


def test_working_calendar():
    cal = Calendar(DEFAULTS)
    h = holidays(2027)
    assert date(2027, 3, 26) in h and date(2027, 3, 29) in h  # Good Friday, Easter Monday 2027
    assert date(2027, 5, 3) in h and date(2027, 10, 4) in h  # QLD Labour Day, King's Birthday
    assert not cal.is_workday(date(2026, 12, 24))  # Christmas shutdown
    assert not cal.is_workday(date(2027, 1, 5))
    assert cal.is_workday(date(2027, 1, 13))
    assert cal.finish(date(2027, 2, 5), 3) == date(2027, 2, 9)  # Fri, Mon, Tue
    assert cal.finish(date(2027, 3, 25), 2) == date(2027, 3, 30)  # Thu, then Easter: Tue


def test_crew_maths():
    c = crew_for(4, DEFAULTS)
    assert (c["plumbers"], c["apprentices"], c["supervisors"], c["effective"]) == (4, 2, 1, 5.0)
    assert crew_for(2, DEFAULTS)["supervisors"] == 0  # a crew of 3 is led by a plumber
    assert plumbers_for(363, 10, DEFAULTS) == 4


@pytest.fixture
def approved_job(client):
    ai.set_reader(FakeReader())
    try:
        r = client.post("/api/quick-start", data={"kind": "hydraulic", "name": "Report job", "client": "Report Co"},
                        files=[("files", ("plans.pdf", _plan_pdf(), "application/pdf"))])
        job = r.json()
    finally:
        ai.set_reader(None)
    assert client.post(f"/api/jobs/{job['id']}/approve-all").status_code == 200
    return job


def _rate(client, **kw):
    r = client.post("/api/productivity", json=kw)
    assert r.status_code == 200, r.text
    return r.json()


def test_hydraulic_reports(client, approved_job):
    job_id = approved_job["id"]
    # No labour rates yet: the hours are missing, not guessed.
    lab = client.get(f"/api/jobs/{job_id}/reports/labour").json()
    assert lab["total_hours"] == 0 and lab["missing"] == 2

    san = _rate(client, kind="hydraulic", service_code="SAN", hours_per_unit=0.5)
    _rate(client, kind="hydraulic", item_type="Floor waste", hours_per_unit=1.0)
    _rate(client, kind="hydraulic", service_code="SAN", dn_min=150, dn_max=300, hours_per_unit=9)  # not DN100: doesn't apply
    lab = client.get(f"/api/jobs/{job_id}/reports/labour").json()
    assert lab["missing"] == 0 and lab["total_hours"] == pytest.approx(18 * 0.5 + 3 * 1.0)
    assert {r["source"] for r in lab["rows"]} == {"productivity rate"}

    mats = client.get(f"/api/jobs/{job_id}/reports/materials").json()["rows"]
    pipe = next(r for r in mats if r["group"] == "Pipe")
    assert pipe["qty"] == pytest.approx(18.0) and pipe["order_qty"] == 3 and pipe["order_uom"] == "× 6 m lengths"
    fw = next(r for r in mats if r["description"] == "Floor waste")
    assert fw["qty"] == 3 and fw["order_qty"] == 3

    # A fixed crew of 2 plumbers (and 1 apprentice) does 12 hours in a day.
    r = client.put(f"/api/jobs/{job_id}/program", json={"start_date": "2027-02-01"})
    p = r.json()
    assert p["mode"] == "crew" and p["stages"][0]["plumbers"] == 2 and p["stages"][0]["apprentices"] == 1
    assert p["stages"][0]["workdays"] == 1 and p["finish_date"] == "2027-02-01"

    # Bigger job with a two-week target: the crew is worked out.
    client.put(f"/api/productivity/{san['id']}", json={"kind": "hydraulic", "service_code": "SAN", "hours_per_unit": 20})
    p = client.put(f"/api/jobs/{job_id}/program", json={"start_date": "2027-02-01", "target_weeks": 2}).json()
    s = p["stages"][0]
    assert p["field_hours"] == pytest.approx(363)
    assert (s["plumbers"], s["apprentices"], s["supervisors"], s["workdays"]) == (4, 2, 1, 10)
    assert p["finish_date"] == "2027-02-12" and p["working_weeks"] == 2 and p["peak_people"] == 7
    assert p["pm_hours"] == pytest.approx(36.3) and p["supervision_hours"] == pytest.approx(80)
    assert [w["week"] for w in p["weekly"]] == ["2027-02-01", "2027-02-08"]

    # Per-stage override: start the stage later with 6 plumbers.
    stage = s["stage"]
    p = client.put(f"/api/jobs/{job_id}/program", json={"start_date": "2027-02-01", "target_weeks": 2,
                                                         "stages": {stage: {"plumbers": 6, "start": "2027-03-01"}}}).json()
    assert p["stages"][0]["start"] == "2027-03-01" and p["stages"][0]["plumbers"] == 6

    x = client.get(f"/api/jobs/{job_id}/reports.xlsx")
    wb = load_workbook(io.BytesIO(x.content))
    assert {"Summary", "Material list", "Labour hours", "Program", "Weekly manpower"} <= set(wb.sheetnames)
    assert wb["Program"]["A2"].value == stage

    # Won tender with a start date shows in the company workforce forecast.
    client.put(f"/api/jobs/{job_id}/tender-result", json={"tender_result": "won", "tender_value": 250000})
    client.put("/api/settings/workforce", json={"headcount": {"plumbers": 3, "apprentices": 2, "supervisors": 1}})
    d = client.get("/api/reports/dashboard?weeks=40").json()
    assert d["by_result"]["won"]["count"] >= 1 and d["win_rate_pct"] is not None
    busy = [w for w in d["workforce"] if w["plumbers"]]
    assert busy and busy[0]["plumbers"] == 6 and busy[0]["short"]["plumbers"] == 3
    assert any(j["name"] == "Report job" for j in d["planned_jobs"])
    assert d["ai"]["plans"] >= 1


def test_civil_labour_and_settings_permissions(client):
    ai.set_reader(FakeCivilReader())
    try:
        job = client.post("/api/quick-start", data={"kind": "civil", "name": "Civil report"},
                          files=[("files", ("civil.pdf", _civil_pdf(), "application/pdf"))]).json()
    finally:
        ai.set_reader(None)
    _rate(client, kind="civil", civil_category="Stormwater drainage", keyword="RCP", unit="m", hours_per_unit=0.5)
    lab = client.get(f"/api/jobs/{job['id']}/reports/labour?include_suggested=true").json()
    pipe = next(r for r in lab["rows"] if "RCP" in r["description"])
    assert pipe["hours"] == pytest.approx(45.5 * 0.5) and lab["missing"] >= 1  # pits and kerb have no rate
    mats = client.get(f"/api/jobs/{job['id']}/reports/materials?include_suggested=true").json()["rows"]
    assert any("not approved" in r["note"] for r in mats)

    # Only admins change company settings; only people with "Edit pricing" change productivity rates.
    client.post("/api/users", json={"email": "noprice@test.local", "password": "password-123", "can_pricing": False})
    tok = client.post("/api/auth/login", json={"email": "noprice@test.local", "password": "password-123"}).json()["token"]
    h = {"Authorization": f"Bearer {tok}"}
    assert client.post("/api/productivity", json={"kind": "civil", "hours_per_unit": 1}, headers=h).status_code == 403
    assert client.get("/api/productivity", headers=h).status_code == 200
    assert client.put("/api/settings/workforce", json={"hours_per_day": 7.6}, headers=h).status_code == 403
    assert client.put("/api/settings/workforce", json={"hours_per_day": 0}).status_code == 422
    assert client.put("/api/settings/workforce", json={"hours_per_day": 7.6}).json()["hours_per_day"] == 7.6
    client.put("/api/settings/workforce", json={"hours_per_day": 8})


def test_productivity_csv_round_trip(client):
    from pathlib import Path

    template = (Path(__file__).resolve().parents[2] / "docs" / "productivity-rates-template.csv").read_bytes()
    r = client.post("/api/productivity/import", files={"file": ("t.csv", template, "text/csv")})
    assert r.status_code == 200 and r.json()["imported"] == 8
    rates = client.get("/api/productivity").json()
    examples = [x for x in rates if "EXAMPLE" in x["note"]]
    assert len(examples) == 8 and not any(x["active"] for x in examples)  # examples never switch themselves on
    csv_text = client.get("/api/productivity.csv").text
    assert csv_text.splitlines()[0].startswith("kind,service_code")
    bad = client.post("/api/productivity/import", files={"file": ("b.csv", b"kind,hours_per_unit\nhydraulic,abc\n", "text/csv")})
    assert bad.status_code == 422
