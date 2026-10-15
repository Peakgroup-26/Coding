"""Tender documents: scopes, specs, acoustic reports and so on are read, and what they say steers the
take-off and the price once a person accepts it. A fake AI reader stands in for the real one."""
import io
import json
import re
import zipfile
from types import SimpleNamespace

import pymupdf
import pytest
from openpyxl import Workbook, load_workbook

from app import ai
from app.ai import TenderItem, TenderReading
from app.models import Job
from app.pricing import loading_for
from app.tender import extract_text, is_document_pdf

from .test_phase7_autopilot import FakeReader, _plan_pdf
from .test_phase7_civil import FakeCivilReader, _civil_pdf

A4 = (595, 842)


def _scope_pdf(text="SCOPE OF WORKS - HYDRAULIC SERVICES") -> bytes:
    doc = pymupdf.open()
    for i in range(2):
        page = doc.new_page(width=A4[0], height=A4[1])
        page.insert_text((60, 80), f"{text} page {i + 1}", fontsize=12)
        page.insert_text((60, 110), "Sanitary drainage in ceilings over apartments shall be acoustic PVC.", fontsize=10)
    data = doc.tobytes()
    doc.close()
    return data


def _docx(paragraphs: list[str]) -> bytes:
    body = "".join(f"<w:p><w:r><w:t>{p}</w:t></w:r></w:p>" for p in paragraphs)
    xml = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?><w:document '
           'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body>' + body + "</w:body></w:document>")
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("word/document.xml", xml)
    return buf.getvalue()


ITEMS = [
    TenderItem("material_spec", "Acoustic drainage pipe in ceilings", "Sanitary drainage in ceilings over apartments shall be acoustic PVC.",
               1, ["SAN"], "", "Acoustic PVC", "", confidence=0.9),
    TenderItem("exclusion", "Fire services by others", "Fire services are by others.", 2, ["FIRE"], confidence=0.95),
    TenderItem("allowance", "Authority fees", "Allow a provisional sum of $5,000 for authority fees.", 2, [],
               description="Provisional sum: authority fees", amount=5000, allowance_kind="provisional", confidence=0.9),
    TenderItem("rate_factor", "Night works", "Connections to the live main are to be made at night.", 2, ["SW"],
               applies_to="labour", confidence=0.8),
    TenderItem("risk", "Liquidated damages", "Liquidated damages of $2,000 per day apply.", 3, confidence=0.9),
]


class TenderFake(FakeReader):
    def __init__(self, items=ITEMS):
        super().__init__()
        self.items = items
        self.contexts = []
        self.docs = []

    def read_sheet(self, page_png, page_text, service_codes, count_types):
        self.contexts.append(getattr(self, "context", ""))
        reading = super().read_sheet(page_png, page_text, service_codes, count_types)
        found = re.search(r"H-\d+", page_text or "")
        reading.sheet_no = found.group(0) if found else reading.sheet_no
        return reading

    def read_tender_doc(self, text, pdf, job_kind, service_codes, materials, categories):
        self.docs.append((text, pdf is not None, job_kind))
        assert "SAN" in service_codes and "PVC DWV" in materials
        return TenderReading("scope", "Hydraulic scope of works", "Acoustic pipe, fire by others.", list(self.items), "fake")


@pytest.fixture
def reader():
    fake = TenderFake()
    ai.set_reader(fake)
    yield fake
    ai.set_reader(None)


def _reqs(client, job_id):
    return {r["title"]: r for r in client.get(f"/api/jobs/{job_id}/requirements").json()["requirements"]}


def test_documents_are_told_apart_from_drawings():
    assert is_document_pdf(_scope_pdf())
    assert not is_document_pdf(_plan_pdf())
    text, pages = extract_text("scope.pdf", _scope_pdf())
    assert pages == 2 and "[Page 2]" in text and "acoustic PVC" in text
    text, _ = extract_text("brief.docx", _docx(["Design brief", "Hot water by heat pump."]))
    assert text == "Design brief\nHot water by heat pump."
    wb = Workbook()
    wb.active.append(["Item", "Note"])
    wb.active.append(["Gas", "Excluded"])
    buf = io.BytesIO()
    wb.save(buf)
    text, _ = extract_text("schedule.xlsx", buf.getvalue())
    assert "Gas\tExcluded" in text
    blank = pymupdf.open()
    blank.new_page(width=A4[0], height=A4[1])
    assert extract_text("scan.pdf", blank.tobytes())[0] == ""  # a scan: no text layer


def test_tender_documents_steer_the_takeoff_and_price(client, reader):
    r = client.post("/api/quick-start", data={"kind": "hydraulic", "name": "Tender job", "client": "Tender Consultants"},
                    files=[("files", ("plans.pdf", _plan_pdf(), "application/pdf")),
                           ("files", ("Scope of works.pdf", _scope_pdf(), "application/pdf"))])
    assert r.status_code == 200, r.text
    job = r.json()
    st = client.get(f"/api/jobs/{job['id']}/autopilot").json()
    assert st["state"] == "done" and st["tender"] == {"documents": 1, "requirements": 5, "suggested": 5, "accepted": 0}
    docs = client.get(f"/api/jobs/{job['id']}/documents").json()
    assert [d["filename"] for d in docs] == ["Scope of works.pdf"] and docs[0]["doc_type"] == "scope"
    assert len(client.get(f"/api/jobs/{job['id']}/files").json()) == 1  # only the plans became sheets
    # The documents were read first, and the drawing reader was told what they say.
    assert "Acoustic PVC" in reader.contexts[0]

    reqs = _reqs(client, job["id"])
    spec = reqs["Acoustic drainage pipe in ceilings"]
    assert spec["page"] == 1 and spec["document"] == "Scope of works.pdf" and "2 runs now" in spec["effect"]
    runs = [x for x in client.get(f"/api/sheets/{client.get(f'/api/jobs/{job['id']}/sheets').json()[0]['id']}/runs").json()]
    assert {x["material"] for x in runs} == {"PVC DWV"}  # nothing changes until it's accepted

    # Accept the spec: the runs change material, and back again on undo.
    r = client.post(f"/api/requirements/{spec['id']}/accept")
    assert r.status_code == 200 and r.json()["status"] == "accepted"
    sheet_id = runs[0]["sheet_id"]
    assert {x["material"] for x in client.get(f"/api/sheets/{sheet_id}/runs").json()} == {"Acoustic PVC"}
    assert "Acoustic PVC" in [m["name"] for m in client.get("/api/settings").json()["materials"]]
    client.post(f"/api/requirements/{spec['id']}/undo")
    assert {x["material"] for x in client.get(f"/api/sheets/{sheet_id}/runs").json()} == {"PVC DWV"}
    client.post(f"/api/requirements/{spec['id']}/accept")

    # The night-works loading needs a % from the estimator before it can be accepted.
    night = reqs["Night works"]
    assert night["needs"] == "Enter the % loading"
    assert client.post(f"/api/requirements/{night['id']}/accept").status_code == 422
    assert client.patch(f"/api/requirements/{night['id']}", json={"percent": 15}).json()["needs"] == ""

    r = client.post(f"/api/jobs/{job['id']}/requirements/accept-all").json()
    assert r == {"accepted": 4, "waiting": 0}
    reqs = client.get(f"/api/jobs/{job['id']}/requirements").json()
    assert reqs["excluded_services"] == ["FIRE"] and reqs["spec_materials"] == {"SAN": "Acoustic PVC"}
    assert reqs["loadings"][0]["percent"] == 15 and reqs["loadings"][0]["service_codes"] == ["SW"]
    allowances = client.get(f"/api/jobs/{job['id']}/allowances").json()
    assert {(a["kind"], a["description"], a["amount"]) for a in allowances} == {
        ("exclusion", "Fire services by others", 0.0), ("provisional", "Provisional sum: authority fees", 5000.0)}

    # A new plan read after accepting picks up the specified material straight away.
    client.post(f"/api/jobs/{job['id']}/add-plans", files=[("files", ("level2.pdf", _plan_pdf("H-301"), "application/pdf"))])
    sheets = client.get(f"/api/jobs/{job['id']}/sheets").json()
    new_sheet = next(s for s in sheets if s["id"] != sheet_id)
    assert {x["material"] for x in client.get(f"/api/sheets/{new_sheet['id']}/runs").json()} == {"Acoustic PVC"}

    # The Excel BOQ lists the tender notes.
    wb = load_workbook(io.BytesIO(client.get(f"/api/jobs/{job['id']}/export.xlsx").content))
    notes = [[c.value for c in row] for row in wb["Tender notes"].iter_rows(min_row=2)]
    assert ["Commercial risk", "Liquidated damages", "Liquidated damages of $2,000 per day apply.", "Scope of works.pdf p3", "Applied"] in notes

    # Removing the document undoes what was accepted from it.
    assert client.delete(f"/api/documents/{docs[0]['id']}").status_code == 204
    reqs = client.get(f"/api/jobs/{job['id']}/requirements").json()
    assert reqs["requirements"] == [] and reqs["excluded_services"] == [] and reqs["spec_materials"] == {}
    assert client.get(f"/api/jobs/{job['id']}/allowances").json() == []


def test_loadings_apply_to_the_services_they_name():
    job = Job(name="x", rate_adjustments=[
        {"label": "Night works", "percent": 15, "applies_to": "labour", "service_codes": ["SW"]},
        {"label": "Remote site", "percent": 5, "applies_to": "all", "service_codes": []}])
    mat, lab, labels = loading_for(job, "SW")
    assert (float(mat), float(lab)) == (5.0, 20.0) and len(labels) == 2
    mat, lab, _ = loading_for(job, "SAN")
    assert (float(mat), float(lab)) == (5.0, 5.0)


def test_civil_tender_loadings_and_exclusions(client):
    class CivilTender(FakeCivilReader):
        def read_tender_doc(self, text, pdf, job_kind, service_codes, materials, categories):
            assert job_kind == "civil" and "Kerb and channel" in categories
            return TenderReading("specification", "Civil spec", "", [
                TenderItem("rate_factor", "Rock excavation", "Rock is expected below 1 m.", 4, [], "Stormwater drainage",
                           percent=10, applies_to="all", confidence=0.8),
                TenderItem("exclusion", "Kerb by others", "Kerb and channel by the roadworks contractor.", 5, [], "Kerb and channel"),
            ], "fake")

    ai.set_reader(CivilTender())
    try:
        client.post("/api/civil/rate-sets", files={"file": ("old.csv", b"Description,Unit,Rate\n375 dia RCP Class 2 stormwater pipe,m,420\n"
                                                                     b"Junction pit 900 x 900,no,2800\n", "text/csv")})
        r = client.post("/api/quick-start", data={"kind": "civil", "name": "Civil tender"},
                        files=[("files", ("civil.pdf", _civil_pdf(), "application/pdf")),
                               ("files", ("spec.docx", _docx(["Rock is expected below 1 m."]), "application/octet-stream"))])
        job = r.json()
        reqs = _reqs(client, job["id"])
        assert "Rock excavation" in reqs and reqs["Rock excavation"]["effect"] == "Add 10% to all costs for Stormwater drainage"
        assert client.post(f"/api/jobs/{job['id']}/requirements/accept-all").json()["accepted"] == 2
        client.post(f"/api/jobs/{job['id']}/autopilot", json={"reread": True})
        boq = client.get(f"/api/jobs/{job['id']}/civil/boq").json()
        assert not [l for l in boq["lines"] if l["category"] == "Kerb and channel"]  # excluded
        client.post(f"/api/jobs/{job['id']}/civil/approve-all")
        s = client.get(f"/api/jobs/{job['id']}/civil/boq").json()["summary"]
        lines = 45.5 * 420 + 3 * 2800
        assert s["lines_total_ex_gst"] == pytest.approx(lines)
        assert s["loadings"][0]["amount"] == pytest.approx(lines * 0.10)
        assert s["total_confirmed_ex_gst"] == pytest.approx(lines * 1.10)
    finally:
        ai.set_reader(None)


def test_claude_reader_reads_tender_documents():
    import anthropic

    from app.ai import ClaudeSheetReader

    reader = ClaudeSheetReader.__new__(ClaudeSheetReader)
    reader.anthropic, reader.model = anthropic, "claude-opus-5-5"
    sent = []
    reply = {"doc_type": "acoustic_report", "title": "Acoustic report", "summary": "Acoustic pipe needed.", "items": [
        {"category": "material_spec", "title": "Acoustic pipe", "quote": "Use acoustic pipe.", "page": -1,
         "service_codes": ["SAN", "XYZ"], "civil_category": "", "material": "Acoustic PVC", "description": "", "qty": -1,
         "unit": "", "amount": -5, "allowance_kind": "", "percent": -2, "applies_to": "", "confidence": 2}]}

    def create(**kw):
        sent.append(kw)
        return SimpleNamespace(stop_reason="end_turn", model="claude-opus-5-5",
                               content=[SimpleNamespace(type="text", text=json.dumps(reply))])

    reader.client = SimpleNamespace(beta=SimpleNamespace(messages=SimpleNamespace(create=create)))
    out = reader.read_tender_doc("", b"%PDF-1.7 scanned", "hydraulic", ["SAN"], ["PVC DWV"], ["Other"])
    item = out.items[0]
    assert out.doc_type == "acoustic_report" and item.service_codes == ["SAN"]
    assert (item.page, item.qty, item.amount, item.percent, item.confidence) == (0, 0, 0, 0, 1.0)
    content = sent[0]["messages"][0]["content"]
    assert content[0]["type"] == "document" and "hydraulic services" in content[1]["text"]

    # The tender context is added to the drawing reads.
    reader.context = "- [material_spec SAN] Acoustic pipe"
    reply = {"sheet_no": "H-1", "title": "", "revision": "", "scale_text": "", "sheet_type": "plan", "services": [],
             "is_hydraulic": True, "confidence": 1, "legend": [], "count_tags": []}
    png = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 20, 20), False).tobytes("png")
    reader.read_sheet(png, "", ["SAN"], ["WC"])
    assert "<tender_requirements>" in sent[-1]["messages"][0]["content"][1]["text"]
