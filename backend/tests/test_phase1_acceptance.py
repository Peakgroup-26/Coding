"""Phase 1 acceptance criteria from the brief:

- A 40-page combined PDF splits correctly.
- Measured lengths on a test line of known length are within ±0.5%.
- Scale verification blocks measuring when the check is out by more than 1%.
- Export totals match on-screen totals exactly.
"""
from io import BytesIO

import pytest
from openpyxl import load_workbook

from .conftest import A3_PT, mm, pdf_bytes


def _job(client, **kw):
    r = client.post("/api/jobs", json={"name": "Test job", "client": "Acme Hydraulics", **kw})
    assert r.status_code == 200, r.text
    return r.json()


def _upload(client, job_id, *files):
    r = client.post(f"/api/jobs/{job_id}/files",
                    files=[("files", (name, data, "application/pdf")) for name, data in files])
    assert r.status_code == 200, r.text
    return r.json()


def _settings(client):
    s = client.get("/api/settings").json()
    return {x["code"]: x for x in s["services"]}, {x["name"]: x for x in s["stages"]}


def _calibrate_and_verify(client, sheet_id, factor=1.0, check_m=5.0):
    r = client.post(f"/api/sheets/{sheet_id}/calibrate", json={
        "method": "two_point", "p1": [mm(100) * factor, mm(100) * factor], "p2": [mm(200) * factor, mm(100) * factor], "known_m": 10.0})
    assert r.status_code == 200, r.text
    r = client.post(f"/api/sheets/{sheet_id}/verify", json={
        "p1": [mm(100) * factor, mm(150) * factor], "p2": [mm(100) * factor, mm(200) * factor], "known_m": check_m})
    assert r.status_code == 200, r.text
    return r.json()


def test_40_page_pdf_splits_and_reads_title_blocks(client):
    job = _job(client)
    pages = [{"sheet_no": f"H-{100 + i}", "title": f"SANITARY DRAINAGE PLAN LEVEL {i}"} for i in range(40)]
    sheets = _upload(client, job["id"], ("combined.pdf", pdf_bytes(pages)))
    assert len(sheets) == 40
    assert [s["page_no"] for s in sheets] == list(range(1, 41))
    assert [s["sheet_no"] for s in sheets] == [f"H-{100 + i}" for i in range(40)]
    assert all(s["revision"] == "A" and s["page_size"] == "A1" for s in sheets)
    assert sheets[0]["stated_scale"] == "1:100 @ A1"
    assert sheets[0]["title"] == "SANITARY DRAINAGE PLAN LEVEL 0"
    assert sheets[0]["classification"] == "plan" and "SAN" in sheets[0]["services_shown"]
    assert sheets[0]["has_vectors"]

    img = client.get(f"/api/sheets/{sheets[0]['id']}/image")
    assert img.status_code == 200 and img.content[:4] == b"\x89PNG"
    snaps = client.get(f"/api/sheets/{sheets[0]['id']}/snap-points").json()
    assert any(abs(x - mm(100)) < 0.5 and abs(y - mm(100)) < 0.5 for x, y in snaps)


def test_many_single_page_files_and_revisions(client):
    job = _job(client)
    _upload(client, job["id"],
            ("h1-b.pdf", pdf_bytes([{"sheet_no": "H-001", "title": "COLD WATER PLAN", "rev": "B"}])),
            ("h1-a.pdf", pdf_bytes([{"sheet_no": "H-001", "title": "COLD WATER PLAN", "rev": "A"}])),
            ("h2.pdf", pdf_bytes([{"sheet_no": "H-002", "title": "HOT WATER PLAN"}])))
    _upload(client, job["id"], ("dup.pdf", pdf_bytes([{"sheet_no": "H-002", "title": "HOT WATER PLAN"}])))

    visible = client.get(f"/api/jobs/{job['id']}/sheets").json()
    everything = client.get(f"/api/jobs/{job['id']}/sheets?include_duplicates=true").json()
    assert len(visible) == 3 and len(everything) == 4
    current = {(s["sheet_no"], s["revision"]) for s in visible if s["is_current"]}
    assert current == {("H-001", "B"), ("H-002", "A")}

    # Measuring on a superseded sheet is refused.
    old = next(s for s in visible if s["revision"] == "A" and s["sheet_no"] == "H-001")
    _calibrate_and_verify(client, old["id"])
    services, _ = _settings(client)
    r = client.post(f"/api/sheets/{old['id']}/runs", json={
        "service_id": services["CW"]["id"], "dn": 20, "material": "Copper", "geometry": [[0, 0], [100, 0]]})
    assert r.status_code == 409


def test_measurement_accuracy_and_scale_blocking(client):
    job = _job(client)
    sheet = _upload(client, job["id"], ("a.pdf", pdf_bytes([{"sheet_no": "H-301", "title": "SANITARY PLAN"}])))[0]
    services, stages = _settings(client)
    run = {"service_id": services["SAN"]["id"], "dn": 100, "material": "PVC DWV",
           "geometry": [[mm(100), mm(100)], [mm(200), mm(100)]]}

    # No scale yet: blocked.
    assert client.post(f"/api/sheets/{sheet['id']}/runs", json=run).status_code == 409

    # Check dimension out by more than 1%: still blocked.
    s = _calibrate_and_verify(client, sheet["id"], check_m=5.2)
    assert s["scale_verified"] is False and s["verification_data"]["error_pct"] > 1
    assert client.post(f"/api/sheets/{sheet['id']}/runs", json=run).status_code == 409

    # Good check: allowed, and the test line measures 10 m within ±0.5%.
    s = _calibrate_and_verify(client, sheet["id"], check_m=5.0)
    assert s["scale_verified"] is True
    r = client.post(f"/api/sheets/{sheet['id']}/runs", json=run)
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["length_m"] == pytest.approx(10.0, rel=0.005)
    assert body["stage_id"] == stages["Under slab"]["id"]  # SAN default stage

    # Recalibrating with runs present is refused, so measured lengths can't silently change.
    r = client.post(f"/api/sheets/{sheet['id']}/calibrate", json={"method": "stated"})
    assert r.status_code == 409


def test_stated_scale_on_a1_drawing_saved_at_a3(client):
    job = _job(client)
    factor = 420 / 841
    sheet = _upload(client, job["id"], ("a3.pdf", pdf_bytes([{
        "sheet_no": "H-401", "title": "STORMWATER PLAN", "size": A3_PT, "factor": factor}])))[0]
    assert sheet["page_size"] == "A3"
    r = client.post(f"/api/sheets/{sheet['id']}/calibrate", json={"method": "stated"})
    assert r.status_code == 200
    r = client.post(f"/api/sheets/{sheet['id']}/verify", json={
        "p1": [mm(100) * factor, mm(150) * factor], "p2": [mm(100) * factor, mm(200) * factor], "known_m": 5.0})
    assert r.json()["scale_verified"] is True

    services, _ = _settings(client)
    r = client.post(f"/api/sheets/{sheet['id']}/runs", json={
        "service_id": services["STW"]["id"], "dn": 150, "material": "PVC SWV",
        "geometry": [[mm(100) * factor, mm(100) * factor], [mm(200) * factor, mm(100) * factor]]})
    assert r.json()["length_m"] == pytest.approx(10.0, rel=0.005)


def test_detail_viewport_has_its_own_scale(client):
    job = _job(client)
    sheet = _upload(client, job["id"], ("d.pdf", pdf_bytes([{"sheet_no": "H-501", "title": "DETAILS"}])))[0]
    _calibrate_and_verify(client, sheet["id"])
    vp = client.post(f"/api/sheets/{sheet['id']}/viewports", json={"name": "Detail 1", "bbox": [mm(300), mm(300), mm(500), mm(450)]}).json()
    services, _ = _settings(client)
    inside = {"service_id": services["CW"]["id"], "dn": 25, "material": "Copper",
              "geometry": [[mm(320), mm(320)], [mm(420), mm(320)]]}
    assert client.post(f"/api/sheets/{sheet['id']}/runs", json=inside).status_code == 409  # viewport not calibrated
    # Detail at 1:20: 100 mm on paper = 2 m.
    client.post(f"/api/viewports/{vp['id']}/calibrate", json={"method": "two_point", "p1": [mm(320), mm(320)], "p2": [mm(420), mm(320)], "known_m": 2.0})
    client.post(f"/api/viewports/{vp['id']}/verify", json={"p1": [mm(320), mm(340)], "p2": [mm(320), mm(390)], "known_m": 1.0})
    r = client.post(f"/api/sheets/{sheet['id']}/runs", json=inside)
    assert r.status_code == 200 and r.json()["length_m"] == pytest.approx(2.0) and r.json()["viewport_id"] == vp["id"]
    crossing = {**inside, "geometry": [[mm(320), mm(320)], [mm(600), mm(320)]]}
    assert client.post(f"/api/sheets/{sheet['id']}/runs", json=crossing).status_code == 409


def test_verticals_waste_counts_and_export_matches_screen(client):
    job = _job(client, quote_no="Q-1001")
    sheet = _upload(client, job["id"], ("e.pdf", pdf_bytes([{"sheet_no": "H-601", "title": "WATER SERVICES PLAN"}])))[0]
    _calibrate_and_verify(client, sheet["id"])
    services, stages = _settings(client)
    jid, sid = job["id"], sheet["id"]

    g = client.post(f"/api/jobs/{jid}/levels", json={"name": "Ground", "rl": 0, "floor_to_floor_m": 3.6}).json()
    l1 = client.post(f"/api/jobs/{jid}/levels", json={"name": "Level 1", "rl": 3.6, "floor_to_floor_m": 3.2}).json()
    client.post(f"/api/jobs/{jid}/levels", json={"name": "Level 2", "rl": 6.8, "floor_to_floor_m": 3.2})
    l2 = client.get(f"/api/jobs/{jid}/levels").json()[2]

    # Admin sets copper waste to 5%.
    mats = {m["name"]: m for m in client.get("/api/settings").json()["materials"]}
    client.put(f"/api/materials/{mats['Copper']['id']}", json={"name": "Copper", "waste_pct": 5, "active": True})

    base = {"service_id": services["CW"]["id"], "material": "Copper"}
    # 10 m horizontal + 3 drops of 0.9 m
    r1 = client.post(f"/api/sheets/{sid}/runs", json={**base, "dn": 25, "geometry": [[mm(100), mm(100)], [mm(200), mm(100)]],
                                                      "vertical_mode": "manual", "vertical_count": 3, "vertical_height_m": 0.9}).json()
    assert r1["vertical_m"] == pytest.approx(2.7) and r1["waste_pct"] == 5
    # Riser Ground to Level 2: 3.6 + 3.2 = 6.8 m, point only
    r2 = client.post(f"/api/sheets/{sid}/runs", json={**base, "dn": 32, "geometry": [[mm(300), mm(300)]],
                                                      "vertical_mode": "levels", "vertical_count": 1,
                                                      "vertical_from_level_id": g["id"], "vertical_to_level_id": l2["id"]}).json()
    assert r2["vertical_m"] == pytest.approx(6.8)
    # A suggested run must not reach the totals.
    client.post(f"/api/sheets/{sid}/runs", json={**base, "dn": 25, "status": "suggested", "geometry": [[mm(100), mm(150)], [mm(100), mm(200)]]})
    # Changing a level height flows through to level-based verticals.
    client.put(f"/api/levels/{l1['id']}", json={"name": "Level 1", "rl": 3.6, "floor_to_floor_m": 3.4})
    assert client.get(f"/api/sheets/{sid}/runs").json()[1]["vertical_m"] == pytest.approx(7.0)

    for _ in range(4):
        client.post(f"/api/sheets/{sid}/counts", json={"service_id": services["SAN"]["id"], "item_type": "WC", "x": 200, "y": 200})
    client.post(f"/api/sheets/{sid}/counts", json={"service_id": services["CW"]["id"], "item_type": "TMV", "x": 210, "y": 200, "status": "suggested"})

    t = client.get(f"/api/jobs/{jid}/totals").json()
    by = {(l["dn"]): l for l in t["lines"]}
    assert by[25]["plan_m"] == pytest.approx(10.0) and by[25]["vertical_m"] == pytest.approx(2.7)
    assert by[25]["waste_m"] == pytest.approx(0.64)  # 5% of 12.7 = 0.635 -> 0.64
    assert by[25]["total_m"] == pytest.approx(13.34)
    assert by[32]["total_m"] == pytest.approx(7.35)  # 7.0 + 0.35
    assert t["grand_total_m"] == pytest.approx(20.69)
    assert t["excluded"] == {"suggested_runs": 1, "suggested_counts": 1, "superseded_runs": 0}
    assert t["counts"] == [{"service_code": "SAN", "item_type": "WC", "stage": "Fit-off", "qty": 4}]

    r = client.get(f"/api/jobs/{jid}/export.xlsx")
    assert r.status_code == 200
    wb = load_workbook(BytesIO(r.content))
    assert {"Summary", "By service", "By stage", "Counts", "Run detail", "Count detail", "Sheet register"} <= set(wb.sheetnames)
    rows = list(wb["By service"].iter_rows(min_row=2, values_only=True))
    assert [(r[0], r[1], r[2], r[3], r[4], r[5], r[6]) for r in rows] == [
        (l["service_code"], l["dn"], l["material"], l["plan_m"], l["vertical_m"], l["waste_m"], l["total_m"]) for l in t["lines"]]
    summary = list(wb["Summary"].iter_rows(values_only=True))
    total_row = next(r for r in summary if r and r[0] == "Total")
    assert total_row[5] == t["grand_total_m"]
    stage_rows = list(wb["By stage"].iter_rows(min_row=2, values_only=True))
    assert sum(r[2] for r in stage_rows) == pytest.approx(t["grand_total_m"])


def test_auth_and_roles(client):
    from fastapi.testclient import TestClient

    from app.main import app
    anon = TestClient(app)
    assert anon.get("/api/jobs").status_code == 401
    assert anon.get("/api/sheets/1/image").status_code == 401

    r = client.post("/api/users", json={"email": "est@test.local", "name": "Est", "role": "estimator", "password": "estimator-pw"})
    assert r.status_code == 200
    tok = anon.post("/api/auth/login", json={"email": "est@test.local", "password": "estimator-pw"}).json()["token"]
    est = {"Authorization": f"Bearer {tok}"}
    assert anon.get("/api/jobs", headers=est).status_code == 200
    assert anon.post("/api/materials", json={"name": "Lead"}, headers=est).status_code == 403
    assert anon.post("/api/auth/login", json={"email": "est@test.local", "password": "wrong"}).status_code == 401
