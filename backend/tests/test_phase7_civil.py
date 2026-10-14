"""Phase 7: civil BOQ from civil plans, priced from previous jobs' rates. A fake AI reader stands in."""
import io

import pytest
from openpyxl import Workbook, load_workbook

from app import ai
from app.ai import CivilGroup, CivilReading, CivilTakeoff
from app.civil import parse_rate_rows, similarity

from .conftest import pdf_bytes


def _rates_xlsx() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.append(["SMITH ST SUBDIVISION - CIVIL BOQ"])
    ws.append([])
    ws.append(["Item", "Description", "Unit", "Qty", "Rate", "Amount"])
    ws.append(["", "STORMWATER DRAINAGE", "", "", "", ""])
    ws.append(["1.1", "375 dia RCP Class 2 stormwater pipe", "lm", 120, "$420.00", 50400])
    ws.append(["1.2", "Junction pit 900 x 900", "no", 4, 2800, 11200])
    ws.append(["1.3", "450 dia RCP Class 2 stormwater pipe", "m", 60, 530, 31800])
    ws.append(["", "ROADWORKS", "", "", "", ""])
    ws.append(["2.1", "Kerb and channel type B", "m", 300, 85, 25500])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


class FakeCivilReader:
    name = "fake"
    model = "fake-model"
    pits = 3

    def read_sheet(self, *a):  # not used for civil jobs
        raise AssertionError

    def read_civil_sheet(self, page_png, page_text, categories):
        assert "Stormwater drainage" in categories
        if "LONG SECTION" in page_text:
            return CivilReading("C-201", "Stormwater long section", "A", "long_section", [
                CivilTakeoff("Stormwater drainage", "375 dia RCP Class 2 pipe SW1-SW3", "m", 45.5, "Long section CH 0-45.5", 0.9)])
        return CivilReading("C-101", "Stormwater plan", "A", "plan", [
            CivilTakeoff("Stormwater drainage", "375 RCP Class 2 pipe, SW1 to SW3", "m", 45.5, "Plan labels", 0.8),
            CivilTakeoff("Stormwater drainage", "Junction pit 900x900", "ea", self.pits, "Pit schedule", 0.95),
            CivilTakeoff("Kerb and channel", "Kerb and channel type B", "m", 0, "Shown on plan, not dimensioned", 0.7, True)])

    def consolidate_civil(self, items):
        pipes = [i["id"] for i in items if "375" in i["description"]]
        rest = [i for i in items if i["id"] not in pipes]
        return ([CivilGroup(pipes, "375 dia RCP Class 2 stormwater pipe", "Stormwater drainage", "m", "max",
                            "Same pipe on plan and long section")]
                + [CivilGroup([i["id"]], i["description"], i["category"], i["unit"], "sum") for i in rest])

    def match_civil_rates(self, lines):
        out = {}
        for l in lines:
            same = [c["rate_id"] for c in l["candidates"]
                    if similarity(l["description"], c["description"]) >= 0.5 and "450" not in c["description"]]
            out[l["line_id"]] = (same, "Same item and size")
        return out


@pytest.fixture
def civil_reader():
    fake = FakeCivilReader()
    ai.set_reader(fake)
    yield fake
    ai.set_reader(None)


def _civil_pdf() -> bytes:
    return pdf_bytes([{"sheet_no": "C-101", "title": "STORMWATER PLAN"}, {"sheet_no": "C-201", "title": "STORMWATER LONG SECTION"}])


def test_rate_import_finds_headers_categories_and_units(client):
    r = client.post("/api/civil/rate-sets", data={"name": "Smith St 2025"},
                    files={"file": ("smith.xlsx", _rates_xlsx(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert r.status_code == 200, r.text
    assert r.json()["rates"] == 4
    rates = client.get(f"/api/civil/rate-sets/{r.json()['id']}/rates").json()
    pipe = rates[0]
    assert pipe["rate"] == 420 and pipe["unit"] == "m" and pipe["category"] == "STORMWATER DRAINAGE" and pipe["code"] == "1.1"
    assert rates[1]["unit"] == "ea" and rates[3]["category"] == "ROADWORKS"

    csv_data = b"Description,Unit,Rate\nTopsoil strip 100mm,m2,4.50\n"
    r = client.post("/api/civil/rate-sets", files={"file": ("old.csv", csv_data, "text/csv")})
    assert r.status_code == 200 and r.json()["rates"] == 1 and r.json()["name"] == "old"
    r = client.post("/api/civil/rate-sets", files={"file": ("x.csv", b"Foo,Bar\n1,2\n", "text/csv")})
    assert r.status_code == 422


def test_similarity_tells_sizes_apart():
    assert similarity("375 RCP Class 2 pipe", "375 dia RCP Class 2 stormwater pipe") > similarity("375 RCP Class 2 pipe", "450 dia RCP Class 2 stormwater pipe")
    rows = parse_rate_rows([["Description", "Rate"], ["Thing", "abc"], ["Other", "12"]])
    assert [r["description"] for r in rows] == ["Other"]


def test_civil_plans_to_priced_boq(client, civil_reader):
    client.post("/api/civil/rate-sets", data={"name": "Smith St 2025"},
                files={"file": ("smith.xlsx", _rates_xlsx(), "application/octet-stream")})
    r = client.post("/api/quick-start", data={"kind": "civil", "name": "Civil job"},
                    files=[("files", ("civil.pdf", _civil_pdf(), "application/pdf"))])
    assert r.status_code == 200, r.text
    job = r.json()
    st = client.get(f"/api/jobs/{job['id']}/autopilot").json()
    assert st["state"] == "done", st

    boq = client.get(f"/api/jobs/{job['id']}/civil/boq").json()
    lines = {l["description"]: l for l in boq["lines"]}
    pipe = lines["375 dia RCP Class 2 stormwater pipe"]
    assert pipe["qty"] == 45.5 and pipe["combine"] == "max" and len(pipe["items"]) == 2  # counted once, not twice
    assert pipe["rate"] == 420 and pipe["rate_sources"][0]["job"] == "Smith St 2025"
    pit = lines["Junction pit 900x900"]
    assert pit["qty"] == 3 and pit["rate"] == 2800
    kerb = lines["Kerb and channel type B"]
    assert kerb["needs_measuring"] and kerb["qty"] == 0
    assert boq["summary"]["lines_confirmed"] == 0 and boq["summary"]["total_confirmed_ex_gst"] == 0

    # Approve: the kerb still needs measuring, so it waits.
    r = client.post(f"/api/jobs/{job['id']}/civil/approve-all").json()
    assert r["approved"] == 2 and r["skipped_need_measuring"] == 1
    assert r["summary"]["total_confirmed_ex_gst"] == pytest.approx(45.5 * 420 + 3 * 2800)

    # Measure the kerb by hand, approve it, and override a rate.
    r = client.patch(f"/api/civil/lines/{kerb['id']}", json={"overridden_qty": 120, "status": "confirmed"})
    assert r.status_code == 200 and r.json()["effective_qty"] == 120 and r.json()["rate"] == 85
    client.patch(f"/api/civil/lines/{pit['id']}", json={"rate": 3000})
    total = 45.5 * 420 + 3 * 3000 + 120 * 85
    assert client.get(f"/api/jobs/{job['id']}/civil/boq").json()["summary"]["total_confirmed_ex_gst"] == pytest.approx(total)

    # Re-reading (say, a new revision with an extra pit) keeps approved lines and hand-typed rates,
    # and flags the changed quantity instead of changing an approved figure.
    civil_reader.pits = 4
    client.post(f"/api/jobs/{job['id']}/autopilot", json={"reread": True})
    boq = client.get(f"/api/jobs/{job['id']}/civil/boq").json()
    assert boq["summary"]["total_confirmed_ex_gst"] == pytest.approx(total)
    assert len([l for l in boq["lines"] if "375" in l["description"]]) == 1
    pit = next(l for l in boq["lines"] if l["id"] == pit["id"])
    assert pit["qty"] == 3 and "now reads 4" in pit["note"]
    assert not next(l for l in boq["lines"] if "375" in l["description"])["note"]

    x = client.get(f"/api/jobs/{job['id']}/civil/boq.xlsx")
    assert x.status_code == 200
    ws = load_workbook(io.BytesIO(x.content)).active
    rows = [[c.value for c in r] for r in ws.iter_rows()]
    assert ["", "TOTAL (ex GST)", "", "", "", pytest.approx(total)] == [v if v is not None else "" for v in rows[-1][:6]]


def test_civil_without_ai_still_takes_manual_lines(client):
    r = client.post("/api/quick-start", data={"kind": "civil"}, files=[("files", ("Bridge_Rd_civil.pdf", _civil_pdf(), "application/pdf"))])
    job = r.json()
    assert job["name"] == "Bridge Rd civil" and job["kind"] == "civil"
    st = client.get(f"/api/jobs/{job['id']}/autopilot").json()
    assert any(a["kind"] == "ai" for a in st["attention"])
    r = client.post(f"/api/jobs/{job['id']}/civil/lines", json={"description": "Junction pit 900 x 900", "unit": "no", "qty": 2})
    assert r.status_code == 200 and r.json()["unit"] == "ea" and r.json()["status"] == "confirmed"
    assert r.json()["rate"] == 2800  # found in the previous job's rates


def test_claude_reader_civil_calls(monkeypatch):
    """The real reader's new calls, with the SDK replaced: request shape and response handling."""
    import json
    from types import SimpleNamespace

    import anthropic
    import pymupdf

    from app.ai import ClaudeSheetReader

    reader = ClaudeSheetReader.__new__(ClaudeSheetReader)
    reader.anthropic, reader.model = anthropic, "claude-opus-5-5"
    replies, sent = [], []

    def create(**kw):
        sent.append(kw)
        return SimpleNamespace(stop_reason="end_turn", model="claude-opus-5-5",
                               content=[SimpleNamespace(type="text", text=json.dumps(replies.pop(0)))])

    reader.client = SimpleNamespace(beta=SimpleNamespace(messages=SimpleNamespace(create=create)))
    png = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 50, 50), False).tobytes("png")

    replies.append({"sheet_no": "C-1", "title": "SW plan", "revision": "A", "sheet_type": "plan", "items": [
        {"category": "Stormwater drainage", "description": "375 RCP", "unit": "m", "quantity": -2, "basis": "label",
         "confidence": 3, "needs_measuring": False}]})
    out = reader.read_civil_sheet(png, "text", ["Stormwater drainage"])
    assert out.items[0].quantity == 0 and out.items[0].confidence == 1.0
    assert sent[-1]["output_config"]["effort"] == "medium"
    assert sent[-1]["output_config"]["format"]["schema"]["properties"]["items"]["items"]["properties"]["category"]["enum"] == ["Stormwater drainage"]

    replies.append({"groups": [{"item_ids": [1, 2], "description": "Pipe", "category": "Stormwater drainage", "unit": "m",
                                "combine": "max", "reason": "same"}]})
    assert reader.consolidate_civil([{"id": 1}, {"id": 2}])[0].combine == "max"
    assert sent[-1]["messages"][0]["content"][0]["type"] == "text"  # text only: no image needed

    replies.append({"matches": [{"line_id": 7, "rate_ids": [3], "reason": "same"}]})
    assert reader.match_civil_rates([{"line_id": 7}]) == {7: ([3], "same")}

    replies.append({"assignments": [
        {"key": "a", "service_code": "SAN", "dn": 100, "material": "PVC DWV", "confidence": 0.9},
        {"key": "b", "service_code": "SAN", "dn": 0, "material": "PVC DWV", "confidence": 0.9}]})
    got = reader.map_styles(png, [], [{"key": "a"}, {"key": "b"}], ["SAN"], ["PVC DWV"])
    assert [g.key for g in got] == ["a"]  # a zero size is dropped
