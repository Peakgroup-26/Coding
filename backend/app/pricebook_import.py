"""Read Peak's price list, prebuilds and labour rates from CSV or Excel.

Column names are matched loosely (case, spacing and punctuation ignored) against common
Simpro export headings, and any column can be named explicitly with a mapping. Prebuilds can
come "long" (one row per component, prebuild code repeated) as Simpro exports them, or "wide"
with a components column like "PART-1 x2; PART-2 x0.5".
"""
import csv
import io
import re
from dataclasses import dataclass, field

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from .models import LabourRate, PriceItem, Prebuild

ALIASES: dict[str, list[str]] = {
    "code": ["code", "part number", "part no", "partno", "part #", "item code", "catalogue code", "catalog code",
             "prebuild code", "prebuild", "sku", "rate code"],
    "description": ["description", "name", "part description", "item description", "prebuild name",
                    "prebuild description", "rate description"],
    "uom": ["uom", "unit", "units", "unit of measure"],
    "cost": ["cost", "cost price", "trade price", "buy price", "unit cost", "price", "cost ex gst", "trade"],
    "supplier": ["supplier", "vendor"],
    "labour_hours": ["labour hours", "labour time", "labor hours", "labour", "hours", "labour hrs", "time", "labour time hrs"],
    "labour_rate_code": ["labour rate", "labour rate code", "rate type"],
    "rate": ["rate", "charge out rate", "chargeout rate", "hourly rate", "sell rate", "rate per hour", "charge rate"],
    "component_code": ["component code", "component", "part code", "material code", "component part number",
                       "catalogue item", "child code"],
    "component_qty": ["component qty", "component quantity", "part qty", "qty", "quantity", "child qty"],
    "components": ["components", "parts"],
}


def _norm(s: str) -> str:
    return re.sub(r"[^a-z0-9#]+", " ", (s or "").lower()).strip()


def read_rows(filename: str, data: bytes) -> list[dict[str, str]]:
    """Rows as {original header: text}, from .csv or .xlsx (first sheet)."""
    if filename.lower().endswith((".xlsx", ".xlsm")):
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        ws = wb.worksheets[0]
        it = ws.iter_rows(values_only=True)
        headers = None
        rows = []
        for row in it:
            if headers is None:
                if row and any(v not in (None, "") for v in row):
                    headers = [str(v).strip() if v is not None else "" for v in row]
                continue
            if not any(v not in (None, "") for v in row):
                continue
            rows.append({h: ("" if v is None else str(v).strip()) for h, v in zip(headers, row) if h})
        return rows
    text = data.decode("utf-8-sig", errors="replace")
    reader = csv.DictReader(io.StringIO(text))
    return [{(k or "").strip(): (v or "").strip() for k, v in r.items() if k} for r in reader]


def resolve_columns(headers: list[str], wanted: list[str], explicit: dict[str, str] | None = None) -> dict[str, str]:
    """Which header holds each wanted field."""
    explicit = explicit or {}
    by_norm = {_norm(h): h for h in headers}
    out = {}
    for f in wanted:
        if explicit.get(f):
            if explicit[f] not in headers:
                raise ValueError(f"Column '{explicit[f]}' (for {f}) isn't in the file")
            out[f] = explicit[f]
            continue
        for alias in ALIASES.get(f, [f]):
            if _norm(alias) in by_norm:
                out[f] = by_norm[_norm(alias)]
                break
    return out


def _num(v: str, default=None):
    v = (v or "").replace("$", "").replace(",", "").strip()
    if v == "":
        return default
    return float(v)


_COMP_RE = re.compile(r"\s*([^;,]+?)\s*(?:[x×*]\s*([0-9.]+))?\s*(?:[;,]|$)")


def parse_components(text: str) -> list[dict]:
    """'PART-1 x2; PART-2 x0.5; PART-3' -> [{code, qty}]"""
    out = []
    for m in _COMP_RE.finditer(text or ""):
        code = m.group(1).strip()
        if code:
            out.append({"code": code, "qty": float(m.group(2)) if m.group(2) else 1.0})
    return out


@dataclass
class ImportResult:
    created: int = 0
    updated: int = 0
    errors: list[dict] = field(default_factory=list)
    columns: dict[str, str] = field(default_factory=dict)

    def as_dict(self):
        return {"created": self.created, "updated": self.updated, "errors": self.errors, "columns": self.columns}


def import_items(db: Session, version_id: int, rows: list[dict], explicit=None) -> ImportResult:
    res = ImportResult()
    if not rows:
        return res
    cols = resolve_columns(list(rows[0].keys()), ["code", "description", "uom", "cost", "supplier"], explicit)
    res.columns = cols
    if "code" not in cols or "cost" not in cols:
        raise ValueError(f"Couldn't find a code and a cost column. Columns found: {', '.join(rows[0].keys())}")
    existing = {p.code: p for p in db.scalars(select(PriceItem).where(PriceItem.version_id == version_id))}
    for i, r in enumerate(rows, start=2):
        code = r.get(cols["code"], "").strip()
        if not code:
            continue
        try:
            cost = _num(r.get(cols["cost"], ""), 0.0)
        except ValueError:
            res.errors.append({"row": i, "code": code, "error": f"Cost '{r.get(cols['cost'])}' isn't a number"})
            continue
        item = existing.get(code)
        if item is None:
            item = PriceItem(version_id=version_id, code=code)
            db.add(item)
            existing[code] = item
            res.created += 1
        else:
            res.updated += 1
        item.cost = cost
        if "description" in cols:
            item.description = r.get(cols["description"], "")
        if "uom" in cols:
            item.uom = r.get(cols["uom"], "") or "ea"
        if "supplier" in cols:
            item.supplier = r.get(cols["supplier"], "")
    db.flush()
    return res


def import_rates(db: Session, version_id: int, rows: list[dict], explicit=None) -> ImportResult:
    res = ImportResult()
    if not rows:
        return res
    cols = resolve_columns(list(rows[0].keys()), ["code", "description", "rate"], explicit)
    res.columns = cols
    if "code" not in cols or "rate" not in cols:
        raise ValueError(f"Couldn't find a code and a rate column. Columns found: {', '.join(rows[0].keys())}")
    existing = {p.code: p for p in db.scalars(select(LabourRate).where(LabourRate.version_id == version_id))}
    for i, r in enumerate(rows, start=2):
        code = r.get(cols["code"], "").strip()
        if not code:
            continue
        try:
            rate = _num(r.get(cols["rate"], ""), 0.0)
        except ValueError:
            res.errors.append({"row": i, "code": code, "error": f"Rate '{r.get(cols['rate'])}' isn't a number"})
            continue
        lr = existing.get(code)
        if lr is None:
            lr = LabourRate(version_id=version_id, code=code)
            db.add(lr)
            existing[code] = lr
            res.created += 1
        else:
            res.updated += 1
        lr.rate = rate
        if "description" in cols:
            lr.description = r.get(cols["description"], "")
    db.flush()
    return res


def import_prebuilds(db: Session, version_id: int, rows: list[dict], explicit=None) -> ImportResult:
    res = ImportResult()
    if not rows:
        return res
    cols = resolve_columns(list(rows[0].keys()),
                           ["code", "description", "uom", "labour_hours", "labour_rate_code",
                            "component_code", "component_qty", "components"], explicit)
    res.columns = cols
    if "code" not in cols:
        raise ValueError(f"Couldn't find a prebuild code column. Columns found: {', '.join(rows[0].keys())}")
    # Group rows by prebuild code (long format), keeping the first non-blank header values.
    groups: dict[str, dict] = {}
    for i, r in enumerate(rows, start=2):
        code = r.get(cols["code"], "").strip()
        if not code:
            continue
        g = groups.setdefault(code, {"row": i, "description": "", "uom": "", "labour_hours": None,
                                     "labour_rate_code": "", "components": []})
        for f in ("description", "uom", "labour_rate_code"):
            if f in cols and not g[f]:
                g[f] = r.get(cols[f], "")
        try:
            if "labour_hours" in cols and g["labour_hours"] is None:
                g["labour_hours"] = _num(r.get(cols["labour_hours"], ""))
            if "component_code" in cols and r.get(cols["component_code"], "").strip():
                qty = _num(r.get(cols["component_qty"], ""), 1.0) if "component_qty" in cols else 1.0
                g["components"].append({"code": r[cols["component_code"]].strip(), "qty": qty})
            if "components" in cols and r.get(cols["components"], "").strip():
                g["components"].extend(parse_components(r[cols["components"]]))
        except ValueError as e:
            res.errors.append({"row": i, "code": code, "error": f"Not a number: {e}"})
    existing = {p.code: p for p in db.scalars(select(Prebuild).where(Prebuild.version_id == version_id))}
    for code, g in groups.items():
        pb = existing.get(code)
        if pb is None:
            pb = Prebuild(version_id=version_id, code=code)
            db.add(pb)
            res.created += 1
        else:
            res.updated += 1
        pb.description = g["description"]
        pb.uom = g["uom"] or "ea"
        pb.labour_hours = g["labour_hours"] or 0.0
        pb.labour_rate_code = g["labour_rate_code"]
        # Merge duplicate component codes
        merged: dict[str, float] = {}
        for c in g["components"]:
            merged[c["code"]] = merged.get(c["code"], 0.0) + c["qty"]
        pb.components = [{"code": k, "qty": v} for k, v in merged.items()]
    db.flush()
    return res
