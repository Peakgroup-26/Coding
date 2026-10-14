"""Civil bill of quantities from a set of civil plans, priced from Peak's previous civil jobs.

1. Each sheet is read by the AI, which lists the quantities written on it (schedules, long
   sections, labelled lengths). It never estimates: an item it can see but can't quantify is
   listed as "needs measuring".
2. The same work often shows on several sheets, so the items are grouped into BOQ lines. The AI
   decides which items are the same work and whether to add them up or count them once; the
   arithmetic is done here.
3. Each line is matched against rates from previous jobs. The suggested rate is the median of
   the matching past rates, and the line keeps a list of which jobs they came from.

Every line starts as a suggestion. Totals and the export only include approved lines.
"""
import csv
import io
import logging
import re
import statistics
from datetime import datetime
from decimal import ROUND_HALF_UP, Decimal

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy import delete, select
from sqlalchemy.orm import Session

from . import storage
from .ai import AIUnavailable
from .audit import record
from .models import CivilBoqLine, CivilRate, CivilRateSet, CivilTakeoffItem, Job, Sheet, User
from .pdf_processing import refresh_currency

CATEGORIES = [
    "Preliminaries", "Site preparation and demolition", "Earthworks", "Stormwater drainage", "Sewer", "Water supply",
    "Subsoil drainage", "Roadworks and pavements", "Kerb and channel", "Footpaths and driveways", "Retaining walls",
    "Erosion and sediment control", "Landscaping", "Electrical and conduits", "Line marking and signage", "Other",
]

UNIT_ALIASES = {
    "m": "m", "lm": "m", "l/m": "m", "lin m": "m", "lin.m": "m", "metre": "m", "metres": "m", "meter": "m", "mtr": "m",
    "m2": "m2", "m²": "m2", "sqm": "m2", "sq m": "m2", "sq.m": "m2",
    "m3": "m3", "m³": "m3", "cum": "m3", "cu m": "m3", "cu.m": "m3",
    "ea": "ea", "each": "ea", "no": "ea", "no.": "ea", "nr": "ea", "number": "ea",
    "item": "item", "ls": "item", "l/s": "item", "lump sum": "item", "sum": "item", "prov": "item", "ps": "item",
    "t": "t", "tonne": "t", "tonnes": "t", "ton": "t",
}

STOP = {"the", "of", "and", "to", "in", "with", "for", "a", "an", "incl", "including", "inc", "supply", "install",
        "lay", "construct", "provide", "all", "as", "per", "on", "at", "mm", "x", "complete", "works", "work"}

TWO_DP = Decimal("0.01")
log = logging.getLogger(__name__)


def norm_unit(u: str) -> str:
    u = (u or "").strip().lower()
    return UNIT_ALIASES.get(u, u)


def tokens(text: str) -> set[str]:
    out = set()
    for t in re.findall(r"[a-z]+|\d+(?:\.\d+)?", (text or "").lower()):
        if t not in STOP:
            out.add(t)
    return out


def similarity(a: str, b: str) -> float:
    ta, tb = tokens(a), tokens(b)
    if not ta or not tb:
        return 0.0
    score = len(ta & tb) / len(ta | tb)
    na, nb = {t for t in ta if t[0].isdigit()}, {t for t in tb if t[0].isdigit()}
    if na and nb and na != nb and not (na & nb):
        score *= 0.5  # different sizes (375 vs 450) are different items
    return score


def money(x) -> Decimal:
    return Decimal(str(x)).quantize(TWO_DP, rounding=ROUND_HALF_UP)


def line_qty(line: CivilBoqLine) -> float:
    return line.overridden_qty if line.overridden_qty is not None else line.qty


# --- importing previous jobs' rates ---
HEADERS = {
    "description": ["description", "item description", "desc", "description of work", "work item", "details", "particulars"],
    "unit": ["unit", "units", "uom", "unit of measure"],
    "rate": ["rate", "unit rate", "rate ($)", "rate $", "$/unit", "unit price", "price", "rate ex gst", "unit rate ($)"],
    "qty": ["qty", "quantity", "quant", "qty."],
    "code": ["item", "item no", "item no.", "item #", "no", "no.", "ref", "code", "item number"],
    "amount": ["amount", "total", "amount ($)", "extended", "value"],
}


def _number(v) -> float | None:
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace("$", "").replace(",", "")
    try:
        return float(s)
    except ValueError:
        return None


def _header_map(row: list) -> dict | None:
    cells = [str(c or "").strip().lower() for c in row]
    found = {}
    for field, names in HEADERS.items():
        for i, c in enumerate(cells):
            if c in names and i not in found.values():
                found[field] = i
                break
    if "description" in found and "rate" in found:
        return found
    return None


def parse_rate_rows(rows: list[list]) -> list[dict]:
    """Rows of a BOQ (any layout with description and rate columns) -> rate dicts.

    Heading rows (text but no rate) set the category for the rows under them.
    """
    out = []
    header, category = None, ""
    for row in rows:
        if header is None:
            header = _header_map(row)
            continue
        def cell(field):
            i = header.get(field)
            return row[i] if i is not None and i < len(row) else None
        desc = str(cell("description") or "").strip()
        rate = _number(cell("rate"))
        if not desc:
            continue
        if rate is None or rate <= 0:
            if not cell("unit") and _number(cell("qty")) is None:
                category = desc[:200]
            continue
        out.append({"code": str(cell("code") or "").strip()[:100], "category": category, "description": desc[:1000],
                    "unit": norm_unit(str(cell("unit") or "")), "rate": rate, "qty": _number(cell("qty"))})
    return out


def read_rate_file(filename: str, data: bytes) -> list[dict]:
    name = filename.lower()
    if name.endswith((".xlsx", ".xlsm")):
        wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
        out = []
        for ws in wb.worksheets:
            out.extend(parse_rate_rows([list(r) for r in ws.iter_rows(values_only=True)]))
        return out
    if name.endswith(".csv"):
        text = data.decode("utf-8-sig", errors="replace")
        return parse_rate_rows([r for r in csv.reader(io.StringIO(text))])
    raise ValueError("Upload rates as an Excel (.xlsx) or CSV file")


# --- reading the plans ---
def run_civil(db: Session, job: Job, user: User, p, reader, reread: bool = False) -> None:
    sheets = list(db.scalars(select(Sheet).where(Sheet.job_id == job.id, Sheet.is_duplicate.is_(False)).order_by(Sheet.file_id, Sheet.page_no)))
    if reader is None or not hasattr(reader, "read_civil_sheet"):
        p.attention("ai", "AI reading is switched off on the server, so civil quantities can't be read automatically. "
                          "Ask your administrator to add the AI key, or add BOQ lines by hand.")
    else:
        for i, sheet in enumerate(sheets):
            if not sheet.is_current:
                continue
            if sheet.ai_reading and not reread:
                continue
            p.step(f"Reading sheet {i + 1} of {len(sheets)}: {sheet.sheet_no or 'p' + str(sheet.page_no)}", i, len(sheets))
            try:
                reading = reader.read_civil_sheet(storage.local_path(sheet.image_key).read_bytes(), sheet.extracted_text, CATEGORIES)
            except AIUnavailable as e:
                p.attention("ai", f"The AI couldn't read sheet {sheet.sheet_no or sheet.page_no}: {e}")
                if "credentials" in str(e) or "reach" in str(e):
                    break
                continue
            except Exception as e:  # one bad sheet shouldn't stop the whole set
                log.exception("Civil AI read failed for sheet %s", sheet.id)
                db.rollback()
                p.attention("ai", f"The AI couldn't read sheet {sheet.sheet_no or sheet.page_no} ({e.__class__.__name__}).")
                continue
            if not sheet.confirmed:
                sheet.sheet_no = reading.sheet_no or sheet.sheet_no
                sheet.title = reading.title or sheet.title
                sheet.revision = reading.revision or sheet.revision
                sheet.classification = reading.sheet_type
                sheet.is_hydraulic = False
            sheet.ai_reading = {"sheet_no": reading.sheet_no, "title": reading.title, "revision": reading.revision,
                                "sheet_type": reading.sheet_type, "model": reading.model, "items": len(reading.items)}
            # The new reading replaces this sheet's earlier quantities. Approved lines keep their approved
            # quantity; build_lines re-attaches the new readings to them and flags any change.
            for old in db.scalars(select(CivilTakeoffItem).where(CivilTakeoffItem.sheet_id == sheet.id)):
                db.delete(old)
            for it in reading.items:
                db.add(CivilTakeoffItem(job_id=job.id, sheet_id=sheet.id, category=it.category, description=it.description[:1000],
                                        unit=norm_unit(it.unit), qty=it.quantity, basis=it.basis, confidence=it.confidence,
                                        needs_measuring=it.needs_measuring or it.quantity == 0))
            db.flush()
            refresh_currency(db, job.id)
            db.commit()
    p.step("Merging items that appear on more than one sheet")
    build_lines(db, job, user, reader, p)
    p.step("Pricing from previous jobs")
    reprice(db, job.id, reader)
    db.commit()
    p.status["summary"] = civil_summary(db, job.id)


def build_lines(db: Session, job: Job, user: User, reader, p=None) -> int:
    """(Re)build the suggested BOQ lines from takeoff items that aren't in a line a person has kept."""
    current = set(db.scalars(select(Sheet.id).where(Sheet.job_id == job.id, Sheet.is_current.is_(True))))
    items = [i for i in db.scalars(select(CivilTakeoffItem).where(CivilTakeoffItem.job_id == job.id).order_by(CivilTakeoffItem.id))
             if i.sheet_id is None or i.sheet_id in current]
    # Old suggestions are rebuilt; anything a person has touched is kept.
    for line in db.scalars(select(CivilBoqLine).where(CivilBoqLine.job_id == job.id, CivilBoqLine.status == "suggested",
                                                      CivilBoqLine.source == "ai")):
        if line.overridden_qty is None and not line.rate_overridden and not line.note:
            db.delete(line)
    db.flush()
    for i in items:
        if i.line_id and db.get(CivilBoqLine, i.line_id) is None:
            i.line_id = None
    items = [i for i in items if i.line_id is None]
    if not items:
        return 0
    sheets = {s.id: s for s in db.scalars(select(Sheet).where(Sheet.job_id == job.id))}
    groups = None
    if reader is not None and hasattr(reader, "consolidate_civil") and len(items) > 1:
        payload = [{"id": i.id, "sheet": (sheets[i.sheet_id].sheet_no if i.sheet_id in sheets else ""),
                    "sheet_title": (sheets[i.sheet_id].title if i.sheet_id in sheets else ""), "category": i.category,
                    "description": i.description, "unit": i.unit, "quantity": i.qty, "basis": i.basis} for i in items]
        try:
            groups = reader.consolidate_civil(payload)
        except Exception as e:  # AIUnavailable or a bad reply: fall back to one line per item
            if p is not None:
                p.attention("ai", f"Couldn't merge repeated items automatically ({e}); each sheet's items are listed separately.")
    by_id = {i.id: i for i in items}
    used: set[int] = set()
    plan = []
    for g in groups or []:
        ids = [x for x in g.item_ids if x in by_id and x not in used]
        units = {by_id[x].unit for x in ids}
        if not ids or len(units) > 1:
            continue  # a group mixing units is a bad merge; those items fall through to their own lines
        used.update(ids)
        plan.append((ids, g.description or by_id[ids[0]].description, g.category or by_id[ids[0]].category,
                     units.pop(), g.combine if g.combine in ("sum", "max") else "sum"))
    for i in items:
        if i.id not in used:
            plan.append(([i.id], i.description, i.category, i.unit, "sum"))
    order = {c: n for n, c in enumerate(CATEGORIES)}
    plan.sort(key=lambda x: (order.get(x[2], len(order)), x[1].lower()))
    kept = [l for l in db.scalars(select(CivilBoqLine).where(CivilBoqLine.job_id == job.id))]
    start = max([l.sort for l in kept], default=0) + 1
    new_lines = 0
    for n, (ids, desc, cat, unit, combine) in enumerate(plan):
        qtys = [by_id[x].qty for x in ids]
        qty = max(qtys) if combine == "max" else sum(qtys)
        same = _matching_kept_line(kept, desc, unit)
        if same is not None:
            # Work a person already approved or edited: attach the new reading and flag any change.
            for x in ids:
                by_id[x].line_id = same.id
            if abs(qty - same.qty) > 1e-6:
                same.note = (f"Plans re-read: quantity now reads {qty:g} {unit} (was {same.qty:g}). "
                             "Check it and update the quantity if needed.")
            continue
        new_lines += 1
        line = CivilBoqLine(job_id=job.id, sort=start + n, category=cat, description=desc[:1000], unit=unit,
                            qty=round(qty, 3), combine=combine,
                            needs_measuring=any(by_id[x].needs_measuring for x in ids) or qty == 0)
        db.add(line)
        db.flush()
        for x in ids:
            by_id[x].line_id = line.id
    record(db, user.id, "job", job.id, "civil_build", after={"job_id": job.id, "lines": new_lines})
    return new_lines


def _matching_kept_line(kept: list[CivilBoqLine], description: str, unit: str) -> CivilBoqLine | None:
    best, best_score = None, 0.6
    for l in kept:
        if norm_unit(l.unit) != norm_unit(unit):
            continue
        s = similarity(l.description, description)
        if s >= best_score:
            best, best_score = l, s
    return best


# --- pricing from previous jobs ---
def candidates(line: CivilBoqLine, rates: list[CivilRate], n: int = 8) -> list[tuple[float, CivilRate]]:
    unit = norm_unit(line.unit)
    scored = []
    for r in rates:
        if unit and r.unit and norm_unit(r.unit) != unit:
            continue
        s = similarity(line.description, r.description)
        if s > 0.15:
            scored.append((s, r))
    scored.sort(key=lambda x: -x[0])
    return scored[:n]


def reprice(db: Session, job_id: int, reader=None) -> int:
    """Suggest a rate for every line that doesn't have one typed in."""
    rates = list(db.scalars(select(CivilRate).join(CivilRateSet).where(CivilRateSet.active.is_(True))))
    sets = {s.id: s.name for s in db.scalars(select(CivilRateSet))}
    lines = [l for l in db.scalars(select(CivilBoqLine).where(CivilBoqLine.job_id == job_id)) if not l.rate_overridden]
    cands = {l.id: candidates(l, rates) for l in lines}
    picked: dict[int, tuple[list[int], str]] = {}
    use_ai = reader is not None and hasattr(reader, "match_civil_rates")
    if use_ai:
        batch = [{"line_id": l.id, "description": l.description, "unit": l.unit,
                  "candidates": [{"rate_id": r.id, "description": r.description, "unit": r.unit, "rate": r.rate,
                                  "job": sets.get(r.set_id, "")} for _, r in cands[l.id]]}
                 for l in lines if cands[l.id]]
        try:
            for k in range(0, len(batch), 40):
                picked.update(reader.match_civil_rates(batch[k:k + 40]))
        except Exception:  # AIUnavailable or a bad reply: fall back to wording matches
            log.warning("AI rate matching failed; using wording matches", exc_info=True)
            use_ai = False
    priced = 0
    for l in lines:
        options = {r.id: r for _, r in cands[l.id]}
        if use_ai:
            ids, why = picked.get(l.id, ([], ""))
            chosen = [options[i] for i in ids if i in options]
        else:
            chosen = [r for s, r in cands[l.id] if s >= 0.5]
            why = "Closest wording in previous jobs" if chosen else ""
        if chosen:
            l.rate = round(statistics.median(r.rate for r in chosen), 2)
            l.rate_sources = [{"rate_id": r.id, "job": sets.get(r.set_id, ""), "description": r.description, "unit": r.unit,
                               "rate": r.rate} for r in chosen]
            l.rate_note = why
            priced += 1
        else:
            l.rate, l.rate_sources = None, []
            l.rate_note = "No similar item in previous jobs" if rates else "No previous job rates loaded yet"
    return priced


# --- totals and export ---
def civil_summary(db: Session, job_id: int) -> dict:
    lines = list(db.scalars(select(CivilBoqLine).where(CivilBoqLine.job_id == job_id)))
    approved = [l for l in lines if l.status == "confirmed"]
    total = sum((money(line_qty(l)) * money(l.rate or 0) for l in approved), Decimal(0))
    return {"lines": len(lines), "lines_suggested": len(lines) - len(approved), "lines_confirmed": len(approved),
            "needs_measuring": sum(1 for l in lines if l.needs_measuring),
            "unpriced": sum(1 for l in lines if l.rate is None),
            "total_confirmed_ex_gst": float(money(total))}


def export_boq(db: Session, job: Job) -> bytes:
    lines = sorted(db.scalars(select(CivilBoqLine).where(CivilBoqLine.job_id == job.id)), key=lambda l: l.sort)
    sheets = {s.id: s for s in db.scalars(select(Sheet).where(Sheet.job_id == job.id))}
    items = list(db.scalars(select(CivilTakeoffItem).where(CivilTakeoffItem.job_id == job.id)))
    wb = Workbook()
    ws = wb.active
    ws.title = "BOQ"
    bold = Font(bold=True)
    ws.append([f"{job.name} - Civil bill of quantities"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([f"Quote {job.quote_no}" if job.quote_no else "", f"Exported {datetime.now():%d/%m/%Y %H:%M}"])
    pending = [l for l in lines if l.status != "confirmed"]
    if pending:
        ws.append([f"{len(pending)} line(s) not yet approved are listed on the 'Not approved' tab and are not in the total."])
        ws["A3"].font = Font(bold=True, color="C00000")
    ws.append([])
    head = ["Item", "Description", "Unit", "Qty", "Rate ($)", "Amount ($)", "From sheets", "Rate from"]
    ws.append(head)
    for c in ws[ws.max_row]:
        c.font = bold
        c.fill = PatternFill("solid", fgColor="DDEBF7")
    grand = Decimal(0)
    approved = [l for l in lines if l.status == "confirmed"]
    n = 0
    for cat in [c for c in CATEGORIES] + sorted({l.category for l in approved} - set(CATEGORIES)):
        rows = [l for l in approved if l.category == cat]
        if not rows:
            continue
        ws.append([cat])
        ws[ws.max_row][0].font = bold
        sub = Decimal(0)
        for l in rows:
            n += 1
            qty = line_qty(l)
            amount = money(qty) * money(l.rate or 0)
            sub += amount
            refs = sorted({sheets[i.sheet_id].sheet_no or f"p{sheets[i.sheet_id].page_no}" for i in items
                           if i.line_id == l.id and i.sheet_id in sheets})
            ws.append([n, l.description, l.unit, qty, float(money(l.rate)) if l.rate is not None else None,
                       float(money(amount)), ", ".join(refs), "; ".join(s["job"] for s in (l.rate_sources or []))])
        ws.append(["", f"Subtotal {cat}", "", "", "", float(money(sub))])
        ws[ws.max_row][1].font = bold
        ws[ws.max_row][5].font = bold
        grand += sub
    ws.append([])
    ws.append(["", "TOTAL (ex GST)", "", "", "", float(money(grand))])
    for c in ws[ws.max_row]:
        c.font = bold
    for col, width in zip("ABCDEFGH", (6, 70, 7, 11, 11, 14, 18, 30)):
        ws.column_dimensions[col].width = width
    for row in ws.iter_rows(min_row=5):
        row[1].alignment = Alignment(wrap_text=True, vertical="top")
        for c in row[3:6]:
            c.number_format = "#,##0.00"

    if pending:
        wp = wb.create_sheet("Not approved")
        wp.append(["Description", "Unit", "Qty", "Suggested rate", "Needs measuring", "Note"])
        for c in wp[1]:
            c.font = bold
        for l in pending:
            wp.append([l.description, l.unit, line_qty(l), l.rate, "Yes" if l.needs_measuring else "", l.rate_note])
    wt = wb.create_sheet("Quantities by sheet")
    wt.append(["Sheet", "Category", "Description", "Unit", "Qty", "Where it came from", "AI confidence"])
    for c in wt[1]:
        c.font = bold
    for i in items:
        s = sheets.get(i.sheet_id)
        wt.append([(s.sheet_no or f"p{s.page_no}") if s else "", i.category, i.description, i.unit, i.qty, i.basis,
                   round(i.confidence, 2)])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def delete_rate_set(db: Session, set_id: int) -> None:
    db.execute(delete(CivilRate).where(CivilRate.set_id == set_id))
    db.execute(delete(CivilRateSet).where(CivilRateSet.id == set_id))
