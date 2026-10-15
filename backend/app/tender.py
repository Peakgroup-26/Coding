"""Tender documents: scopes, specifications, acoustic reports, briefs, addenda and conditions.

The AI reads each document and lists what changes the take-off or the price, quoting the
document's own words and page. Each item is a suggestion; accepting it applies its effect:

  material_spec  the service's pipe material changes on its runs, and the autopilot uses it from
                 then on. The price follows, because price mappings go by material.
  exclusion      the service is left out of the take-off (its suggestions are removed) and an
                 exclusion is added to the estimate.
  allowance      an allowance (PC item, provisional sum or item) is added to the estimate, or a
                 line to a civil BOQ.
  rate_factor    a % loading on labour, material or both, for the services or civil category it
                 names. The estimator sets the % unless the document states one.
  inclusion, risk, note
                 recorded and listed in the export's Tender notes; no automatic effect.

What accepting changed is stored on the requirement so it can be undone.
"""
import io
import logging
import re
import zipfile
from datetime import datetime, timezone
from xml.etree import ElementTree

import pymupdf
from sqlalchemy import select
from sqlalchemy.orm import Session

from . import storage
from .ai import AIUnavailable
from .audit import record, snapshot
from .models import (CivilBoqLine, CountItem, Job, JobAllowance, Material, Run, Service, TenderDocument, TenderRequirement,
                     User)
from .scale import detect_page_size

log = logging.getLogger(__name__)

DOC_EXTENSIONS = (".pdf", ".docx", ".xlsx", ".xlsm", ".txt", ".csv", ".md")
CHUNK_CHARS = 150_000  # per AI call; long specifications are read in parts
MAX_SCANNED_PDF_BYTES = 30 * 1024 * 1024
ACTIONABLE = ("material_spec", "exclusion", "allowance", "rate_factor")


# --- telling documents from drawings, and getting their text ---
def is_document_pdf(data: bytes) -> bool:
    """A PDF that is mostly A4 pages is a document (scope, report, spec), not a drawing set."""
    try:
        doc = pymupdf.open(stream=data, filetype="pdf")
    except Exception:
        return False
    try:
        pages = [doc[i] for i in range(min(doc.page_count, 6))]
        a4 = sum(1 for p in pages if detect_page_size(p.rect.width, p.rect.height) == "A4"
                 or max(p.rect.width, p.rect.height) < 900)  # also US letter and other small pages
        return bool(pages) and a4 / len(pages) >= 0.5
    finally:
        doc.close()


def extract_text(filename: str, data: bytes) -> tuple[str, int]:
    """Text with [Page n] markers, and the page count. Empty text means a scanned PDF."""
    name = filename.lower()
    if name.endswith(".pdf"):
        doc = pymupdf.open(stream=data, filetype="pdf")
        try:
            parts = []
            for i, page in enumerate(doc):
                t = (page.get_text("text") or "").strip()
                if t:
                    parts.append(f"[Page {i + 1}]\n{t}")
            text = "\n\n".join(parts)
            if len(text) < 10 * doc.page_count:  # next to no text layer: a scan
                text = ""
            return text, doc.page_count
        finally:
            doc.close()
    if name.endswith(".docx"):
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            root = ElementTree.fromstring(z.read("word/document.xml"))
        ns = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
        paras = []
        for para in root.iter(f"{ns}p"):
            t = "".join(x.text or "" for x in para.iter(f"{ns}t")).strip()
            if t:
                paras.append(t)
        return "\n".join(paras), 0
    if name.endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
        parts = []
        for ws in wb.worksheets:
            rows = ["\t".join("" if c is None else str(c) for c in r) for r in ws.iter_rows(values_only=True)]
            parts.append(f"[Sheet {ws.title}]\n" + "\n".join(r for r in rows if r.strip()))
        return "\n\n".join(parts), 0
    return data.decode("utf-8-sig", errors="replace"), 0


def _chunks(text: str) -> list[str]:
    if len(text) <= CHUNK_CHARS:
        return [text]
    out, start = [], 0
    while start < len(text):
        end = min(len(text), start + CHUNK_CHARS)
        cut = text.rfind("\n[Page ", start, end) if end < len(text) else end
        if cut <= start:
            cut = end
        out.append(text[start:cut])
        start = cut
    return out


# --- reading ---
def read_documents(db: Session, job: Job, user: User, p, reader) -> None:
    """Read every queued tender document on the job and add its requirements as suggestions."""
    from .civil import CATEGORIES

    docs = list(db.scalars(select(TenderDocument).where(TenderDocument.job_id == job.id, TenderDocument.status == "queued")))
    if not docs:
        return
    if reader is None or not hasattr(reader, "read_tender_doc"):
        for d in docs:
            d.text, d.page_count = extract_text(d.filename, storage.read_bytes(d.storage_key))
            d.status = "read"
            d.summary = "Not read: AI reading is switched off on the server."
        db.commit()
        p.attention("ai", "Tender documents were stored but not read, because AI reading is switched off.")
        return
    services = [s.code for s in db.scalars(select(Service).where(Service.active.is_(True)))]
    materials = [m.name for m in db.scalars(select(Material).where(Material.active.is_(True)))]
    for n, d in enumerate(docs):
        p.step(f"Reading tender document {n + 1} of {len(docs)}: {d.filename}", n, len(docs))
        try:
            data = storage.read_bytes(d.storage_key)
            d.text, d.page_count = extract_text(d.filename, data)
            readings = []
            if d.text:
                for part in _chunks(d.text):
                    readings.append(reader.read_tender_doc(part, None, job.kind, services, materials, CATEGORIES))
            elif d.filename.lower().endswith(".pdf") and len(data) <= MAX_SCANNED_PDF_BYTES:
                readings.append(reader.read_tender_doc("", data, job.kind, services, materials, CATEGORIES))
            else:
                raise ValueError("no text could be read from it (a scanned file over 30 MB, or an empty file)")
        except (AIUnavailable, ValueError) as e:
            db.rollback()
            d = db.get(TenderDocument, d.id)
            d.status, d.error = "failed", str(e)
            db.commit()
            p.attention("documents", f"{d.filename}: couldn't be read ({e}).")
            continue
        except Exception as e:  # one bad document shouldn't stop the rest
            log.exception("Tender document %s failed", d.id)
            db.rollback()
            d = db.get(TenderDocument, d.id)
            d.status, d.error = "failed", e.__class__.__name__
            db.commit()
            p.attention("documents", f"{d.filename}: couldn't be read ({e.__class__.__name__}).")
            continue
        first = readings[0]
        d.doc_type, d.title = first.doc_type, first.title[:500]
        d.summary = " ".join(r.summary for r in readings if r.summary)
        d.status, d.error = "read", ""
        seen = set()
        for r in readings:
            for it in r.items:
                key = (it.category, it.title.lower(), it.quote[:80].lower())
                if key in seen:
                    continue
                seen.add(key)
                db.add(TenderRequirement(
                    job_id=job.id, document_id=d.id, category=it.category, title=it.title[:300], quote=it.quote,
                    page=it.page or None, service_codes=it.service_codes, civil_category=it.civil_category,
                    material=it.material[:100], description=it.description, qty=it.qty, unit=it.unit, amount=it.amount,
                    allowance_kind=it.allowance_kind, percent=it.percent, applies_to=it.applies_to, confidence=it.confidence))
        record(db, user.id, "tender_document", d.id, "ai_read", after={"job_id": job.id, "type": d.doc_type})
        db.commit()


def context_for(db: Session, job_id: int) -> str:
    """A short summary of the job's tender requirements, for the AI reading the drawings."""
    lines = []
    for r in db.scalars(select(TenderRequirement).where(TenderRequirement.job_id == job_id,
                                                        TenderRequirement.status != "rejected",
                                                        TenderRequirement.category.in_(("material_spec", "exclusion", "inclusion", "allowance")))):
        svc = ",".join(r.service_codes or [])
        extra = f" material: {r.material}" if r.material else ""
        lines.append(f"- [{r.category}{' ' + svc if svc else ''}] {r.title}.{extra} \"{r.quote[:300]}\"")
    return "\n".join(lines[:80])


# --- applying ---
def _services(db: Session) -> dict[str, Service]:
    return {s.code: s for s in db.scalars(select(Service))}


def missing_input(req: TenderRequirement) -> str:
    """Why a requirement can't be accepted yet, or ''."""
    if req.category == "material_spec" and (not req.material.strip() or not req.service_codes):
        return "Choose the service and the material"
    if req.category == "exclusion" and not req.service_codes and not req.civil_category:
        return "Choose what's excluded"
    if req.category == "rate_factor" and req.percent <= 0:
        return "Enter the % loading"
    return ""


def effect(db: Session, req: TenderRequirement) -> str:
    """What accepting this will do, in plain words."""
    job = db.get(Job, req.job_id)
    svcs = ", ".join(req.service_codes or []) or "all services"
    if req.category == "material_spec":
        n = db.query(Run).filter(Run.job_id == job.id, Run.service_id.in_(
            [s.id for c, s in _services(db).items() if c in (req.service_codes or [])])).count()
        return f"Change {svcs} pipe to {req.material or '?'} ({n} run{'s' if n != 1 else ''} now) and use it from now on"
    if req.category == "exclusion":
        if job.kind == "civil":
            return f"Leave {req.civil_category} out of the BOQ and list it as an exclusion"
        return f"Leave {svcs} out of the take-off and list it as an exclusion"
    if req.category == "allowance":
        what = f"${req.amount:,.2f}" if req.amount else "an amount you enter"
        if job.kind == "civil":
            return f"Add a BOQ line: {req.description or req.title} ({what})"
        return f"Add a {({'pc': 'PC item', 'provisional': 'provisional sum'}).get(req.allowance_kind, 'allowance')} of {what}"
    if req.category == "rate_factor":
        on = req.applies_to or "labour"
        target = (req.civil_category or "all BOQ lines") if job.kind == "civil" else svcs
        return f"Add {req.percent:g}% to {on} costs for {target}" if req.percent else "Add a % loading (enter the %)"
    return "Note it in the tender notes"


def apply(db: Session, req: TenderRequirement, user: User) -> None:
    """Accept a requirement and make its change. Raises ValueError if it needs more input."""
    why = missing_input(req)
    if why:
        raise ValueError(why)
    job = db.get(Job, req.job_id)
    services = _services(db)
    applied: dict = {}
    if req.category == "material_spec":
        name = req.material.strip()
        mat = db.scalar(select(Material).where(Material.name == name))
        if mat is None:
            mat = Material(name=name, waste_pct=0.0)
            db.add(mat)
            db.flush()
            applied["created_material"] = name
        spec = dict(job.spec_materials or {})
        applied["previous_spec"] = {c: spec.get(c) for c in req.service_codes}
        for c in req.service_codes:
            spec[c] = name
        job.spec_materials = spec
        ids = [services[c].id for c in req.service_codes if c in services]
        changed = {}
        for run in db.scalars(select(Run).where(Run.job_id == job.id, Run.service_id.in_(ids))):
            if run.material != name:
                changed[str(run.id)] = [run.material, run.waste_pct]
                before = snapshot(run)
                run.material, run.waste_pct = name, mat.waste_pct
                record(db, user.id, "run", run.id, "update", before, snapshot(run))
        applied["runs"] = changed
    elif req.category == "exclusion":
        if req.service_codes:
            excl = list(job.excluded_services or [])
            applied["added"] = [c for c in req.service_codes if c not in excl]
            job.excluded_services = excl + applied["added"]
            ids = [services[c].id for c in req.service_codes if c in services]
            removed = 0
            for model in (Run, CountItem):
                for obj in db.scalars(select(model).where(model.job_id == job.id, model.service_id.in_(ids), model.status == "suggested")):
                    db.delete(obj)
                    removed += 1
            applied["removed_suggestions"] = removed
        if job.kind == "civil" and req.civil_category:
            excl = list(job.excluded_services or [])
            key = f"civil:{req.civil_category}"
            applied["added"] = applied.get("added", []) + ([key] if key not in excl else [])
            job.excluded_services = excl + [k for k in [key] if k not in excl]
            removed = 0
            for line in db.scalars(select(CivilBoqLine).where(CivilBoqLine.job_id == job.id, CivilBoqLine.category == req.civil_category,
                                                              CivilBoqLine.status == "suggested", CivilBoqLine.source == "ai")):
                db.delete(line)
                removed += 1
            applied["removed_suggestions"] = applied.get("removed_suggestions", 0) + removed
        if job.kind != "civil":
            svc = services.get(req.service_codes[0]) if req.service_codes else None
            a = JobAllowance(job_id=job.id, kind="exclusion", description=req.title[:500],
                             service_id=svc.id if svc else None)
            db.add(a)
            db.flush()
            applied["allowance_id"] = a.id
    elif req.category == "allowance":
        if job.kind == "civil":
            last = max([l.sort for l in db.scalars(select(CivilBoqLine).where(CivilBoqLine.job_id == job.id))], default=0)
            has_amount = req.amount > 0
            line = CivilBoqLine(job_id=job.id, sort=last + 1, category=req.civil_category or "Preliminaries",
                                description=(req.description or req.title)[:1000],
                                unit="item" if has_amount else (req.unit or "item"),
                                qty=1 if has_amount else req.qty, rate=req.amount if has_amount else None,
                                rate_overridden=has_amount, rate_note="From the tender documents" if has_amount else "",
                                needs_measuring=not has_amount and not req.qty,
                                status="confirmed" if has_amount else "suggested", source="manual",
                                note=f"From tender document: {req.quote[:300]}")
            db.add(line)
            db.flush()
            applied["civil_line_id"] = line.id
        else:
            svc = services.get(req.service_codes[0]) if req.service_codes else None
            kind = req.allowance_kind if req.allowance_kind in ("pc", "provisional") else "provisional"
            a = JobAllowance(job_id=job.id, kind=kind, description=(req.description or req.title)[:500],
                             amount=req.amount, service_id=svc.id if svc else None,
                             stage_id=svc.default_stage_id if svc else None)
            db.add(a)
            db.flush()
            applied["allowance_id"] = a.id
    elif req.category == "rate_factor":
        adj = list(job.rate_adjustments or [])
        adj.append({"requirement_id": req.id, "label": req.title, "percent": req.percent,
                    "applies_to": req.applies_to or "labour", "service_codes": req.service_codes or [],
                    "civil_category": req.civil_category})
        job.rate_adjustments = adj
    before = snapshot(req)
    req.status, req.applied = "accepted", applied
    req.decided_by, req.decided_at = user.id, datetime.now(timezone.utc)
    record(db, user.id, "tender_requirement", req.id, "accept", before, snapshot(req))


def undo(db: Session, req: TenderRequirement, user: User, new_status: str = "suggested") -> None:
    """Reverse what accepting did (as far as it can be) and set the status back."""
    job = db.get(Job, req.job_id)
    a = req.applied or {}
    if req.status == "accepted":
        if req.category == "material_spec":
            spec = dict(job.spec_materials or {})
            for c, prev in (a.get("previous_spec") or {}).items():
                if prev:
                    spec[c] = prev
                else:
                    spec.pop(c, None)
            job.spec_materials = spec
            for run_id, (material, waste) in (a.get("runs") or {}).items():
                run = db.get(Run, int(run_id))
                if run is not None and run.material == req.material.strip():
                    before = snapshot(run)
                    run.material, run.waste_pct = material, waste
                    record(db, user.id, "run", run.id, "update", before, snapshot(run))
        elif req.category == "exclusion":
            job.excluded_services = [c for c in (job.excluded_services or []) if c not in (a.get("added") or [])]
        elif req.category == "rate_factor":
            job.rate_adjustments = [x for x in (job.rate_adjustments or []) if x.get("requirement_id") != req.id]
        if a.get("allowance_id"):
            obj = db.get(JobAllowance, a["allowance_id"])
            if obj is not None:
                db.delete(obj)
        if a.get("civil_line_id"):
            obj = db.get(CivilBoqLine, a["civil_line_id"])
            if obj is not None:
                db.delete(obj)
    before = snapshot(req)
    req.status, req.applied = new_status, None
    req.decided_by, req.decided_at = user.id, datetime.now(timezone.utc)
    record(db, user.id, "tender_requirement", req.id, "undo" if new_status == "suggested" else new_status, before, snapshot(req))


# --- notes for exports ---
CATEGORY_LABELS = {
    "material_spec": "Material specified", "exclusion": "Exclusion", "inclusion": "Inclusion", "allowance": "Allowance",
    "rate_factor": "Rate loading", "risk": "Commercial risk", "note": "Note",
}


def notes_rows(db: Session, job_id: int) -> list[list]:
    docs = {d.id: d for d in db.scalars(select(TenderDocument).where(TenderDocument.job_id == job_id))}
    rows = []
    for r in db.scalars(select(TenderRequirement).where(TenderRequirement.job_id == job_id).order_by(TenderRequirement.category, TenderRequirement.id)):
        if r.status == "rejected":
            continue
        d = docs.get(r.document_id)
        src = f"{d.filename}{f' p{r.page}' if r.page else ''}" if d else ""
        rows.append([CATEGORY_LABELS.get(r.category, r.category), r.title, r.quote, src,
                     "Applied" if r.status == "accepted" else "Not yet accepted"])
    return rows


def safe_name(filename: str) -> str:
    return re.sub(r"[^A-Za-z0-9._-]", "_", filename)[-120:]
