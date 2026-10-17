"""Daily to-do lists.

Everyone can keep their own list for the day. Supervisors, foremen, managers and the office can set lists for the
team (a manager for the people who report to them). Before someone clocks off, every task on their list for that
day has to be ticked off, or marked "couldn't finish" with a reason, which goes back to whoever set it.
"""
import json
import re
from datetime import date, datetime, timedelta, timezone

from sqlalchemy import select
from sqlalchemy.orm import Session

from . import features
from . import timesheets as ts
from .models import DailyTask, EmployeeProfile, Job, User
from .plant import notify

LEADS = ("supervisor", "foreman", "project manager", "manager")


def _name(u: User | None) -> str:
    return (u.name or u.email) if u else ""


def leads_team(db: Session, user: User) -> bool:
    """Can set lists for other people (not just their own)."""
    if user.role == "admin" or user.trade in LEADS:
        return True
    if user.role != "field" and features.level(db, user, "todos") == "edit":
        return True
    return db.scalar(select(EmployeeProfile.id).where(EmployeeProfile.manager_id == user.id).limit(1)) is not None


def can_set_for(db: Session, user: User, target_id: int) -> bool:
    if target_id == user.id:
        return True
    if user.role == "admin" or user.trade in LEADS or (user.role != "field" and features.level(db, user, "todos") == "edit"):
        return True
    p = db.scalar(select(EmployeeProfile).where(EmployeeProfile.user_id == target_id))
    return bool(p and p.manager_id == user.id)


def team(db: Session, user: User) -> list[User]:
    """The people this person can set lists for, themselves first."""
    people = [u for u in db.scalars(select(User).where(User.active.is_(True)).order_by(User.name)) if u.id != user.id and can_set_for(db, user, u.id)]
    return [user] + people


def today() -> date:
    return ts.local_day(ts.now())


def work_day(db: Session, user: User) -> date:
    """The day whose list counts: the day they clocked on (a night shift finishing after midnight still counts yesterday)."""
    cur = ts.open_entry(db, user.id)
    return ts.local_day(cur.start_at) if cur else today()


def task_out(t: DailyTask, users: dict[int, User], jobs: dict[int, Job] | None = None) -> dict:
    j = (jobs or {}).get(t.job_id or 0)
    return {"id": t.id, "user_id": t.user_id, "user": _name(users.get(t.user_id)), "day": t.day.isoformat(), "text": t.text, "job_id": t.job_id,
            "job": f"{j.job_no or ''} {j.name}".strip() if j else "", "set_by_id": t.set_by, "set_by": _name(users.get(t.set_by or 0)),
            "status": t.status, "comment": t.comment, "done_at": t.done_at.isoformat() if t.done_at else None, "carried_from": t.carried_from}


def lists(db: Session, user_ids: list[int], day: date) -> list[DailyTask]:
    if not user_ids:
        return []
    return list(db.scalars(select(DailyTask).where(DailyTask.user_id.in_(user_ids), DailyTask.day == day)
                           .order_by(DailyTask.user_id, DailyTask.position, DailyTask.id)))


def add(db: Session, actor: User, user_ids: list[int], texts: list[str], day: date, job_id: int | None = None) -> list[DailyTask]:
    out = []
    for uid in dict.fromkeys(user_ids):
        if not can_set_for(db, actor, uid):
            raise PermissionError("You can only set to-dos for yourself and the people you look after")
        pos = len(lists(db, [uid], day))
        made = []
        for text in texts:
            text = text.strip()
            if not text:
                continue
            t = DailyTask(user_id=uid, day=day, text=text[:500], job_id=job_id, set_by=actor.id, position=pos)
            pos += 1
            db.add(t)
            made.append(t)
        db.flush()
        if made and uid != actor.id:
            when = "today" if day == today() else "tomorrow" if day == today() + timedelta(days=1) else day.strftime("%a %d/%m")
            what = made[0].text if len(made) == 1 else f"{len(made)} to-dos"
            notify(db, uid, "todo", f"{_name(actor)} set {what} for you {when}.", "/field", key=f"todo-{made[0].id}")
        out += made
    return out


def set_status(db: Session, actor: User, t: DailyTask, status: str, comment: str = "") -> DailyTask:
    if actor.id != t.user_id and not can_set_for(db, actor, t.user_id):
        raise PermissionError("That's not on your list")
    comment = comment.strip()
    if status == "not_done" and len(comment) < 3:
        raise ValueError("Say why it couldn't be finished")
    t.status = status
    t.comment = comment if status in ("done", "not_done") else ""  # a note on a finished task is optional
    t.done_at = datetime.now(timezone.utc) if status != "open" else None
    if status == "not_done" and t.set_by and t.set_by != actor.id:
        notify(db, t.set_by, "todo", f"{_name(actor)} couldn't finish \"{t.text}\": {comment}", "/todos", key=f"todo-nd-{t.id}-{int(t.done_at.timestamp())}")
    return t


def unresolved(db: Session, user: User, day: date | None = None) -> list[DailyTask]:
    """Tasks still open: neither ticked off nor explained."""
    day = day or work_day(db, user)
    return [t for t in lists(db, [user.id], day) if t.status == "open"]


def clock_off_problems(db: Session, user: User) -> list[str]:
    left = unresolved(db, user)
    if not left:
        return []
    names = ", ".join(f"\"{t.text}\"" for t in left[:3]) + (f" and {len(left) - 3} more" if len(left) > 3 else "")
    return [f"Before you clock off, tick off your to-dos or say why you couldn't finish them: {names}."]


def carry_over(db: Session, actor: User, from_day: date, to_day: date, user_ids: list[int]) -> int:
    """Put yesterday's unfinished tasks (not done, or never touched) on the new day's list."""
    n = 0
    already = {t.carried_from for t in lists(db, user_ids, to_day) if t.carried_from}
    for t in lists(db, user_ids, from_day):
        if t.status == "done" or t.id in already or not can_set_for(db, actor, t.user_id):
            continue
        db.add(DailyTask(user_id=t.user_id, day=to_day, text=t.text, job_id=t.job_id, set_by=t.set_by or actor.id, carried_from=t.id,
                         position=100 + t.position))
        n += 1
    db.flush()
    return n


# --- talk a list in: "Dave, pressure test level 2 and clean the slab. Sam, pack down the shield. Everyone lock up." ---
PARSE_PROMPT = """You turn a supervisor's spoken instructions into daily to-do lists for a plumbing company's crew.
Today is {today} ({weekday}). The speaker is {me} (user id {me_id}).

The people the speaker can give tasks to (id: name, trade):
{people}

Live jobs (id: number and name) a task may belong to:
{jobs}

What they said (speech-to-text, so expect missing punctuation and misheard names):
\"\"\"{text}\"\"\"

Work out who gets which tasks. Rules:
- A name (or a close mishearing of one) starts that person's tasks until another name comes up.
- "Everyone", "the whole crew", "all the boys", "you all" means every person marked "crew" except the speaker.
- "I", "me", "myself" means the speaker.
- Split joined instructions into separate tasks ("pressure test level 2 and clean up the slab" is two tasks).
- Write each task as a short instruction starting with a capital letter, in Australian English, keeping their wording,
  measurements and job references. Drop filler ("um", "can you", "make sure you").
- Set job_id only when the job is clearly named; otherwise null.
- If they say the list is for tomorrow or a weekday, give that date as "day"; otherwise null.
- Anything you can't give to a person goes in "unmatched".

Reply with JSON only, no other text:
{{"day": "YYYY-MM-DD" or null, "lists": [{{"user_id": 1, "tasks": [{{"text": "...", "job_id": null}}]}}], "unmatched": ["..."]}}"""

_FILLER = re.compile(r"^(?:(?:and|then|also|so|ok(?:ay)?|right|um+|uh+|can you|could you|i need you to|i want you to|you need to|need to|"
                     r"please|make sure (?:you )?|just|today|you(?:'re| are)? (?:going to|gonna)|go and|you)\s*[,:]?\s+)+", re.I)


def on_tools(people: list[User], me: User) -> list[User]:
    """'Everyone' when a supervisor is talking: the crew on the tools, not the office."""
    return [u for u in people if u.id != me.id and (u.role == "field" or u.trade in ("plumber", "apprentice", "labourer", "foreman", "supervisor"))]


def _clean(task: str) -> str:
    t = _FILLER.sub("", task.strip(" ,.;:-")).strip(" ,.;:-")
    t = re.sub(r"\s+(?:today|tomorrow|this arvo|this afternoon|this morning|please|thanks|cheers)$", "", t, flags=re.I).strip(" ,.;:-")
    return (t[:1].upper() + t[1:])[:500] if t else ""


_PREPS = {"off", "from", "to", "with", "for", "tell", "ask", "call", "ring", "text", "give", "help", "by", "at", "about", "pick", "drop", "meet", "see"}
_GROUP = re.compile(r"\b(?:the whole crew|whole crew|all of you|you all|you guys|all the boys|the boys|the crew|everybody|everyone)\b", re.I)


def _heuristic(text: str, me: User, people: list[User], day: date) -> dict:
    """No AI: read word by word. A name (not straight after 'off', 'from', 'tell'...) starts that person's tasks;
    'everyone' starts the crew's, 'I' at the start of a sentence starts the speaker's; 'and' / 'then' / commas split tasks."""
    firsts: dict[str, User] = {}
    for u in people:
        first = (u.name or u.email).strip().split()[0].lower()
        firsts.setdefault(first, u)
    crew = on_tools(people, me)
    if re.search(r"\btomorrow\b", text, re.I):
        day = day + timedelta(days=1)
    text = _GROUP.sub(" EVERYONE ", text)
    segs: list[tuple[list[User], list[str]]] = []
    names: list[User] = []
    words: list[str] = []
    prev = ""
    sentence_start = True
    for tok in text.split():
        w = tok.strip(".,;:!?").lower()
        new: list[User] | None = None
        if tok.strip(".,;:!?") == "EVERYONE":
            new = crew
        elif w in firsts and prev not in _PREPS:
            new = [firsts[w]]
        elif sentence_start and w in ("i", "i'll", "i'm", "me", "myself"):
            new = [me]
        if new is not None:
            if words or not names:
                if words:
                    segs.append((names, words))
                names, words = list(new), []
            else:
                names += [u for u in new if u not in names]  # "Dave and Sam, ..."
            if tok[-1:] in ".!?;":
                words.append(".")
        elif w in ("and", "&") and names and not words:
            pass
        else:
            words.append(tok)
        prev = w
        sentence_start = tok[-1:] in ".!?;"
    if words:
        segs.append((names, words))
    lists: dict[int, list[str]] = {}
    unmatched: list[str] = []
    last: list[User] = []
    for who, ws in segs:
        who = who or last
        last = who
        body = " ".join(ws)
        tasks = []
        for part in re.split(r"[.;!?\n]+|,\s*|\s+(?:and then|then|and also|and|also)\s+", body):
            t = _clean(re.sub(r"^(?:i(?:'ll| will| need to| have to| gotta)?|'ll)\s+", "", part.strip(), flags=re.I))
            if t and not re.fullmatch(r"(?i)(today|tomorrow|thanks|cheers|that's it|that's all|please)", t):
                tasks.append(t)
        if not who:
            unmatched += tasks
            continue
        for u in who:
            have = lists.setdefault(u.id, [])
            have += [t for t in tasks if t not in have]
    return {"day": day, "lists": [{"user_id": uid, "tasks": [{"text": t, "job_id": None} for t in ts_]} for uid, ts_ in lists.items() if ts_],
            "unmatched": unmatched}


def parse_spoken(db: Session, me: User, text: str, day: date | None = None) -> dict:
    """Turn what a supervisor said into a list per person, for them to check before it's saved."""
    from . import chat

    day = day or today()
    people = team(db, me)
    ids = {u.id: u for u in people}
    crew = on_tools(people, me)
    jobs = list(db.scalars(select(Job).where(Job.work_status.in_(("scheduled", "in_progress", "on_hold"))).order_by(Job.id.desc()).limit(80)))
    out, used_ai = None, False
    model = chat.get_model()
    if model is not None:
        prompt = PARSE_PROMPT.format(today=day.isoformat(), weekday=day.strftime("%A"), me=_name(me), me_id=me.id, text=text[:6000],
                                     people="\n".join(f"{u.id}: {_name(u)}, {u.trade or 'staff'}{', crew' if u in crew else ''}" for u in people),
                                     jobs="\n".join(f"{j.id}: {j.job_no or ''} {j.name}".strip() for j in jobs) or "(none)")
        try:
            raw = model.summarise(prompt)
            m = re.search(r"\{.*\}", raw, re.S)
            data = json.loads(m.group(0)) if m else None
            if isinstance(data, dict) and isinstance(data.get("lists"), list):
                d = day
                try:
                    d = date.fromisoformat(data["day"]) if data.get("day") else day
                except (TypeError, ValueError):
                    pass
                out = {"day": d, "lists": [], "unmatched": [str(x) for x in data.get("unmatched") or []]}
                for li in data["lists"]:
                    uid = li.get("user_id") if isinstance(li, dict) else None
                    if uid not in ids:
                        out["unmatched"] += [str(t.get("text", t)) if isinstance(t, dict) else str(t) for t in (li or {}).get("tasks", [])]
                        continue
                    tasks = []
                    for t in li.get("tasks") or []:
                        txt = _clean(str(t.get("text", "")) if isinstance(t, dict) else str(t))
                        jid = t.get("job_id") if isinstance(t, dict) else None
                        if txt:
                            tasks.append({"text": txt, "job_id": jid if any(j.id == jid for j in jobs) else None})
                    if tasks:
                        out["lists"].append({"user_id": uid, "tasks": tasks})
                used_ai = True
        except Exception:  # noqa: BLE001 - fall back to the simple reading rather than lose what they said
            out = None
    if out is None:
        out = _heuristic(text, me, people, day)
    jobs_by = {j.id: j for j in jobs}
    merged: dict[int, list[dict]] = {}
    for li in out["lists"]:
        merged.setdefault(li["user_id"], []).extend(li["tasks"])
    return {"day": out["day"].isoformat(), "used_ai": used_ai, "unmatched": out["unmatched"],
            "lists": [{"user_id": uid, "name": _name(ids[uid]), "is_me": uid == me.id,
                       "tasks": [{**t, "job": f"{jobs_by[t['job_id']].job_no or ''} {jobs_by[t['job_id']].name}".strip() if t.get("job_id") in jobs_by else ""}
                                 for t in tasks]}
                      for uid, tasks in merged.items()]}


# --- who gets through their list: the management report ---
def _period_key(d: date, group: str) -> str:
    if group == "month":
        return d.strftime("%Y-%m")
    return (d - timedelta(days=d.weekday())).isoformat()  # week starting Monday


def rating(done_pct: float | None, untouched: int, n: int) -> str:
    if not n:
        return "No to-dos"
    if done_pct is None:
        return "No to-dos"
    if done_pct >= 85 and untouched <= n * 0.05:
        return "Performing"
    if done_pct >= 65:
        return "Watch"
    return "Needs attention"


def report(db: Session, viewer: User, start: date, end: date, group: str = "week") -> dict:
    """Per person: how many to-dos they got, finished, couldn't finish (with reasons) or never touched, by week or month."""
    t0 = today()
    people = team(db, viewer)
    ids = [u.id for u in people]
    tasks = list(db.scalars(select(DailyTask).where(DailyTask.user_id.in_(ids), DailyTask.day >= start, DailyTask.day <= min(end, t0))
                            .order_by(DailyTask.day, DailyTask.id))) if ids else []
    # today's still-open tasks aren't late yet: leave them out until the day is over
    tasks = [t for t in tasks if not (t.day == t0 and t.status == "open")]
    periods = sorted({_period_key(t.day, group) for t in tasks} | {_period_key(min(end, t0), group), _period_key(start, group)})
    by_user: dict[int, list[DailyTask]] = {}
    for t in tasks:
        by_user.setdefault(t.user_id, []).append(t)
    rows = []
    for u in people:
        ts_ = by_user.get(u.id, [])
        n = len(ts_)
        done = sum(1 for t in ts_ if t.status == "done")
        nd = sum(1 for t in ts_ if t.status == "not_done")
        untouched = n - done - nd
        days = {}
        for t in ts_:
            days.setdefault(t.day, []).append(t)
        per = {}
        for p in periods:
            pts = [t for t in ts_ if _period_key(t.day, group) == p]
            per[p] = {"tasks": len(pts), "done": sum(1 for t in pts if t.status == "done"),
                      "done_pct": round(100 * sum(1 for t in pts if t.status == "done") / len(pts), 1) if pts else None}
        done_pct = round(100 * done / n, 1) if n else None
        rows.append({"user_id": u.id, "name": _name(u), "trade": u.trade, "tasks": n, "done": done, "not_done": nd, "untouched": untouched,
                     "done_pct": done_pct, "actioned_pct": round(100 * (done + nd) / n, 1) if n else None,
                     "days": len(days), "days_all_done": sum(1 for v in days.values() if all(t.status == "done" for t in v)),
                     "set_by_others": sum(1 for t in ts_ if t.set_by and t.set_by != u.id),
                     "carried": sum(1 for t in ts_ if t.carried_from), "periods": per, "rating": rating(done_pct, untouched, n),
                     "reasons": [{"day": t.day.isoformat(), "task": t.text, "reason": t.comment} for t in ts_ if t.status == "not_done"][-20:]})
    order = {"Needs attention": 0, "Watch": 1, "Performing": 2, "No to-dos": 3}
    rows.sort(key=lambda r: (order[r["rating"]], r["done_pct"] if r["done_pct"] is not None else 101, r["name"]))
    n_all = sum(r["tasks"] for r in rows)
    return {"start": start.isoformat(), "end": end.isoformat(), "group": group, "periods": periods, "people": rows,
            "totals": {"tasks": n_all, "done": sum(r["done"] for r in rows), "not_done": sum(r["not_done"] for r in rows),
                       "untouched": sum(r["untouched"] for r in rows),
                       "done_pct": round(100 * sum(r["done"] for r in rows) / n_all, 1) if n_all else None},
            "thresholds": "Performing: 85%+ finished, with no more than 1 in 20 never actioned. Watch: 65–85%. Needs attention: under 65%."}


def report_xlsx(rep: dict) -> bytes:
    import io

    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    label = (lambda p: datetime.strptime(p, "%Y-%m").strftime("%b %Y")) if rep["group"] == "month" else \
        (lambda p: "w/c " + date.fromisoformat(p).strftime("%d/%m/%Y"))
    fills = {"Performing": "D6F0DE", "Watch": "FFF1D9", "Needs attention": "FDE5E5"}
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws.append([f"To-do performance {date.fromisoformat(rep['start']):%d/%m/%Y} to {date.fromisoformat(rep['end']):%d/%m/%Y}"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([rep["thresholds"]])
    ws.append([])
    head = ["Person", "Trade", "Rating", "To-dos", "Finished", "Couldn't finish", "Never actioned", "% finished", "% actioned", "Days with a list",
            "Days all finished", "Set by others", "Carried over"]
    ws.append(head)
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    for r in rep["people"]:
        ws.append([r["name"], r["trade"], r["rating"], r["tasks"], r["done"], r["not_done"], r["untouched"], r["done_pct"], r["actioned_pct"], r["days"],
                   r["days_all_done"], r["set_by_others"], r["carried"]])
        if r["rating"] in fills:
            ws.cell(row=ws.max_row, column=3).fill = PatternFill("solid", fgColor=fills[r["rating"]])
    for i, w in enumerate([24, 14, 16, 9, 10, 15, 15, 11, 11, 15, 16, 13, 12], start=1):
        ws.column_dimensions[ws.cell(row=4, column=i).column_letter].width = w
    ws2 = wb.create_sheet("By " + rep["group"])
    ws2.append(["Person"] + [label(p) for p in rep["periods"]])
    for c in ws2[1]:
        c.font = Font(bold=True)
    for r in rep["people"]:
        ws2.append([r["name"]] + [(f"{r['periods'][p]['done_pct']}% ({r['periods'][p]['done']}/{r['periods'][p]['tasks']})"
                                   if r["periods"][p]["tasks"] else "") for p in rep["periods"]])
    ws2.column_dimensions["A"].width = 24
    ws3 = wb.create_sheet("Reasons")
    ws3.append(["Person", "Day", "Task", "Why it couldn't be finished"])
    for c in ws3[1]:
        c.font = Font(bold=True)
    for r in rep["people"]:
        for x in r["reasons"]:
            ws3.append([r["name"], date.fromisoformat(x["day"]).strftime("%d/%m/%Y"), x["task"], x["reason"]])
    for col, w in zip("ABCD", (24, 12, 50, 60)):
        ws3.column_dimensions[col].width = w
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
