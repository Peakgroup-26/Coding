"""One Excel workbook with every report for a job: summary, material list, labour hours, crew and
program (with a Gantt chart in the cells), weekly manpower and cash flow."""
import io
from datetime import date, timedelta

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from . import reports
from .models import Job

BOLD = Font(bold=True)
HEAD = PatternFill("solid", fgColor="DDEBF7")
BAR = PatternFill("solid", fgColor="2F6FD6")
WEEKEND = PatternFill("solid", fgColor="F2F2F2")


def _table(ws, headers: list[str], rows: list[list], widths: list[int] | None = None) -> None:
    ws.append(headers)
    for c in ws[ws.max_row]:
        c.font, c.fill = BOLD, HEAD
    for r in rows:
        ws.append(r)
    for i, w in enumerate(widths or [], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def workbook(db: Session, job: Job, include_suggested: bool = False) -> bytes:
    mats = reports.materials(db, job, include_suggested, by_stage=True)
    lab = reports.labour(db, job, include_suggested)
    prog = reports.program(db, job, include_suggested)
    wb = Workbook()

    ws = wb.active
    ws.title = "Summary"
    ws.append([f"{job.name}: reports"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append(["Quantities", "approved and suggested" if include_suggested else "approved only"])
    ws.append([])
    for label, value in [
        ("Labour (field hours)", prog["field_hours"]), ("Items with no labour rate", lab["missing"]),
        ("Supervision hours", prog["supervision_hours"]), ("Project management hours", prog["pm_hours"]),
        ("Plumbers (typical crew)", prog["base_plumbers"]), ("Peak people on site", prog["peak_people"]),
        ("Start", prog["start_date"]), ("Finish", prog["finish_date"]), ("Working weeks", prog["working_weeks"]),
        ("Calendar weeks", prog["calendar_weeks"]),
    ]:
        ws.append([label, value])
    ws.append([])
    ws.append(["Assumptions (admin settings)"])
    ws[ws.max_row][0].font = BOLD
    for k, v in prog["settings"].items():
        ws.append([k.replace("_", " "), v])
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 34, 22

    wm = wb.create_sheet("Material list")
    _table(wm, ["Stage", "Group", "Service", "Code", "Description", "Qty", "Unit", "Order", "Order unit", "Note"],
           [[r["stage"], r["group"], r.get("service_code", ""), r.get("code", ""), r["description"], r["qty"], r["uom"],
             r["order_qty"], r["order_uom"], r.get("note", "")] for r in mats["rows"]], [16, 20, 8, 14, 46, 10, 6, 8, 16, 18])

    wl = wb.create_sheet("Labour hours")
    _table(wl, ["Stage", "Service", "Type", "Description", "Qty", "Unit", "Hours / unit", "Hours", "From"],
           [[r["stage"], r["service_code"], r["kind"], r["description"], r["qty"], r["uom"], r["hours_per_unit"], r["hours"], r["source"]]
            for r in lab["rows"]], [16, 8, 8, 46, 10, 6, 12, 10, 18])
    wl.append([])
    wl.append(["Total", "", "", "", "", "", "", lab["total_hours"]])
    wl[wl.max_row][0].font = BOLD

    wp = wb.create_sheet("Program")
    headers = ["Stage", "Hours", "Plumbers", "Apprentices", "Supervisors", "Working days", "Start", "Finish"]
    _table(wp, headers, [], [18, 9, 9, 11, 11, 12, 11, 11])
    if prog["stages"]:
        first = date.fromisoformat(prog["start_date"])
        first -= timedelta(days=first.weekday())
        last = date.fromisoformat(prog["finish_date"])
        weeks = []
        w = first
        while w <= last:
            weeks.append(w)
            w += timedelta(weeks=1)
        col0 = len(headers) + 1
        for i, wk in enumerate(weeks):
            c = wp.cell(row=1, column=col0 + i, value=wk.strftime("%d/%m"))
            c.font, c.fill = BOLD, HEAD
            c.alignment = Alignment(text_rotation=90)
            wp.column_dimensions[get_column_letter(col0 + i)].width = 3.2
        for s in prog["stages"]:
            wp.append([s["stage"], s["hours"], s["plumbers"], s["apprentices"], s["supervisors"], s["workdays"],
                       date.fromisoformat(s["start"]).strftime("%d/%m/%Y"), date.fromisoformat(s["end"]).strftime("%d/%m/%Y")])
            row = wp.max_row
            st, en = date.fromisoformat(s["start"]), date.fromisoformat(s["end"])
            for i, wk in enumerate(weeks):
                if wk <= en and wk + timedelta(days=6) >= st:
                    wp.cell(row=row, column=col0 + i).fill = BAR

    ww = wb.create_sheet("Weekly manpower")
    _table(ww, ["Week starting", "Plumbers", "Apprentices", "Supervisors", "People", "Field hours", "Stages"],
           [[date.fromisoformat(w["week"]).strftime("%d/%m/%Y"), w["plumbers"], w["apprentices"], w["supervisors"],
             w["plumbers"] + w["apprentices"] + w["supervisors"], w["hours"], ", ".join(w["stages"])] for w in prog["weekly"]],
           [14, 10, 12, 12, 9, 12, 50])

    if prog["cashflow"]:
        wc = wb.create_sheet("Cash flow")
        _table(wc, ["Month", "Value earned (ex GST)", "Cumulative"],
               [[c["month"], c["value"], c["cumulative"]] for c in prog["cashflow"]], [12, 22, 16])
        for row in wc.iter_rows(min_row=2, min_col=2, max_col=3):
            for c in row:
                c.number_format = "#,##0.00"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
