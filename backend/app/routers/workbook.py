"""Estimating workbook (CostX / CANDY): resources and build-ups, workbook rows and variables, take-off dimensions,
subcontract packages and quote levelling, estimate versions and the tender programme / cash flow."""
import csv
import io
import mimetypes
import uuid
from datetime import date, datetime, timezone

from fastapi import APIRouter, Depends, File, Form, HTTPException, Response, UploadFile
from pydantic import BaseModel, Field
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from .. import config, storage
from .. import estimating as est
from ..access import can_see_job
from ..audit import record
from ..auth import current_user
from ..db import get_db
from ..models import (EstimateLine, EstimateSnapshot, EstimateVariable, Job, PackageQuote, RateBuildup, Resource, TenderProgramme,
                      TradePackage, User)
from . import get_or_404
from .estimate import ensure_job_open

router = APIRouter(prefix="/api", tags=["workbook"])


def _job(db: Session, job_id: int, user: User, edit: bool = False) -> Job:
    j = get_or_404(db, Job, job_id)
    if not can_see_job(user, j):
        raise HTTPException(404, "Job not found")
    if edit:
        ensure_job_open(j)
    return j


# --- resources ---
class ResourceIn(BaseModel):
    code: str = Field(min_length=1, max_length=50)
    kind: str = Field(default="material", pattern="^(labour|plant|material|subcontract|other)$")
    description: str = ""
    unit: str = "ea"
    rate: float = Field(default=0, ge=0)
    supplier: str = ""


def resource_out(r: Resource, used: dict[str, int] | None = None) -> dict:
    return {"id": r.id, "code": r.code, "kind": r.kind, "description": r.description, "unit": r.unit, "rate": r.rate, "supplier": r.supplier,
            "updated_at": r.updated_at.isoformat() if r.updated_at else None, "used_in": (used or {}).get(r.code, 0)}


@router.get("/resources")
def list_resources(db: Session = Depends(get_db), _: User = Depends(current_user)):
    used: dict[str, int] = {}
    for b in db.scalars(select(RateBuildup)):
        for ln in b.lines or []:
            used[ln.get("resource", "")] = used.get(ln.get("resource", ""), 0) + 1
    return [resource_out(r, used) for r in db.scalars(select(Resource).order_by(Resource.kind, Resource.code))]


@router.post("/resources")
def add_resource(body: ResourceIn, db: Session = Depends(get_db), _: User = Depends(current_user)):
    if db.scalar(select(Resource).where(func.lower(Resource.code) == body.code.strip().lower())):
        raise HTTPException(409, "That code is taken")
    r = Resource(**{**body.model_dump(), "code": body.code.strip()})
    db.add(r)
    db.commit()
    return resource_out(r)


@router.put("/resources/{res_id}")
def edit_resource(res_id: int, body: ResourceIn, db: Session = Depends(get_db), user: User = Depends(current_user)):
    r = get_or_404(db, Resource, res_id)
    before = r.rate
    old_code = r.code
    for k, v in body.model_dump().items():
        setattr(r, k, v)
    r.updated_at = datetime.now(timezone.utc)
    if old_code != r.code:  # keep build-ups and rows pointing at it
        for b in db.scalars(select(RateBuildup)):
            if any(ln.get("resource") == old_code for ln in b.lines or []):
                b.lines = [{**ln, "resource": r.code} if ln.get("resource") == old_code else ln for ln in b.lines]
        for ln in db.scalars(select(EstimateLine).where(EstimateLine.rate_source == "resource", EstimateLine.rate_code == old_code)):
            ln.rate_code = r.code
    if before != r.rate:
        record(db, user.id, "resource", r.id, "rate", before={"rate": before}, after={"rate": r.rate})
    db.commit()
    return resource_out(r)


@router.delete("/resources/{res_id}", status_code=204)
def delete_resource(res_id: int, db: Session = Depends(get_db), _: User = Depends(current_user)):
    r = get_or_404(db, Resource, res_id)
    for b in db.scalars(select(RateBuildup)):
        if any(ln.get("resource") == r.code for ln in b.lines or []):
            raise HTTPException(409, f"Build-up {b.code} uses it. Take it out of the build-up first.")
    db.delete(r)
    db.commit()


@router.post("/resources/import")
async def import_resources(file: UploadFile = File(...), db: Session = Depends(get_db), _: User = Depends(current_user)):
    """CSV or Excel with columns code, kind, description, unit, rate (and supplier). Existing codes are updated."""
    data = await file.read(config.MAX_UPLOAD_BYTES + 1)
    rows: list[dict] = []
    name = (file.filename or "").lower()
    if name.endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook

        ws = load_workbook(io.BytesIO(data), read_only=True, data_only=True).active
        it = ws.iter_rows(values_only=True)
        head = [str(h or "").strip().lower() for h in next(it, [])]
        rows = [dict(zip(head, r)) for r in it]
    else:
        text = data.decode("utf-8-sig", errors="replace")
        rows = [{(k or "").strip().lower(): v for k, v in r.items()} for r in csv.DictReader(io.StringIO(text))]
    added = updated = 0
    skipped = []
    for i, row in enumerate(rows, start=2):
        code = str(row.get("code") or "").strip()
        if not code:
            continue
        try:
            rate = float(str(row.get("rate") or 0).replace("$", "").replace(",", ""))
        except ValueError:
            skipped.append(f"row {i}: rate isn't a number")
            continue
        kind = str(row.get("kind") or "material").strip().lower()
        if kind not in est.RESOURCE_KINDS:
            kind = "material"
        r = db.scalar(select(Resource).where(Resource.code == code))
        if r is None:
            r = Resource(code=code)
            db.add(r)
            added += 1
        else:
            updated += 1
        r.kind, r.rate = kind, rate
        r.description = str(row.get("description") or r.description or "")
        r.unit = str(row.get("unit") or r.unit or "ea")
        r.supplier = str(row.get("supplier") or r.supplier or "")
        r.updated_at = datetime.now(timezone.utc)
    db.commit()
    return {"added": added, "updated": updated, "skipped": skipped}


# --- rate build-ups ---
class BuildupIn(BaseModel):
    code: str = Field(min_length=1, max_length=50)
    description: str = ""
    unit: str = "m"
    lines: list[dict] = []
    notes: str = ""


@router.get("/rate-buildups")
def list_buildups(db: Session = Depends(get_db), _: User = Depends(current_user)):
    res = est.resource_map(db)
    return [{"id": b.id, "code": b.code, "description": b.description, "unit": b.unit, "lines": b.lines, "notes": b.notes,
             "rate": est.buildup_unit(b, res)} for b in db.scalars(select(RateBuildup).order_by(RateBuildup.code))]


def _check_buildup(db: Session, body: BuildupIn) -> list[dict]:
    res = est.resource_map(db)
    out = []
    for ln in body.lines:
        code = str(ln.get("resource") or "").strip()
        if code not in res:
            raise HTTPException(422, f"No resource with code {code}")
        out.append({"resource": code, "qty": float(ln.get("qty") or 0), "waste_pct": float(ln.get("waste_pct") or 0)})
    return out


@router.post("/rate-buildups")
def add_buildup(body: BuildupIn, db: Session = Depends(get_db), _: User = Depends(current_user)):
    if db.scalar(select(RateBuildup).where(func.lower(RateBuildup.code) == body.code.strip().lower())):
        raise HTTPException(409, "That code is taken")
    b = RateBuildup(code=body.code.strip(), description=body.description, unit=body.unit, lines=_check_buildup(db, body), notes=body.notes)
    db.add(b)
    db.commit()
    return {"id": b.id, "rate": est.buildup_unit(b, est.resource_map(db))}


@router.put("/rate-buildups/{bu_id}")
def edit_buildup(bu_id: int, body: BuildupIn, db: Session = Depends(get_db), _: User = Depends(current_user)):
    b = get_or_404(db, RateBuildup, bu_id)
    old = b.code
    b.code, b.description, b.unit, b.lines, b.notes = body.code.strip(), body.description, body.unit, _check_buildup(db, body), body.notes
    if old != b.code:
        for ln in db.scalars(select(EstimateLine).where(EstimateLine.rate_source == "buildup", EstimateLine.rate_code == old)):
            ln.rate_code = b.code
    db.commit()
    return {"id": b.id, "rate": est.buildup_unit(b, est.resource_map(db))}


@router.delete("/rate-buildups/{bu_id}", status_code=204)
def delete_buildup(bu_id: int, db: Session = Depends(get_db), _: User = Depends(current_user)):
    b = get_or_404(db, RateBuildup, bu_id)
    if db.scalar(select(EstimateLine.id).where(EstimateLine.rate_source == "buildup", EstimateLine.rate_code == b.code).limit(1)):
        raise HTTPException(409, "Workbook rows use it")
    db.delete(b)
    db.commit()


# --- the workbook ---
@router.get("/jobs/{job_id}/workbook")
def get_workbook(job_id: int, db: Session = Depends(get_db), user: User = Depends(current_user)):
    return est.workbook(db, _job(db, job_id, user))


@router.get("/jobs/{job_id}/dimensions")
def get_dimensions(job_id: int, db: Session = Depends(get_db), user: User = Depends(current_user)):
    return sorted(est.dimensions(db, _job(db, job_id, user)).values(), key=lambda d: d["key"])


class LineIn(BaseModel):
    section: str = Field(default="General", min_length=1, max_length=100)
    position: float | None = None
    ref: str = ""
    description: str = Field(min_length=1, max_length=500)
    qty_source: str = Field(default="manual", pattern="^(manual|formula|takeoff)$")
    qty: float = 0
    formula: str = ""
    dimension: str = ""
    uom: str = "item"
    rate_source: str = Field(default="manual", pattern="^(manual|buildup|resource)$")
    rate_code: str = ""
    labour_hours: float = Field(default=0, ge=0)
    labour_rate: float = Field(default=0, ge=0)
    material: float = 0
    plant: float = 0
    subcontract: float = 0
    margin_pct: float | None = None
    service_id: int | None = None
    stage_id: int | None = None
    note: str = ""


def _validate_line(db: Session, body: LineIn) -> None:
    from ..expr import FormulaError, evaluate

    if body.qty_source in ("formula",) and not body.formula.strip():
        raise HTTPException(422, "Type the formula, e.g. weeks * 1")
    if body.qty_source == "takeoff" and not body.dimension:
        raise HTTPException(422, "Pick the take-off quantity to link to")
    if body.formula:
        try:
            evaluate(body.formula, q=1, **{k: 1 for k in ("weeks", "levels")})
        except FormulaError as e:
            if not str(e).startswith("Unknown name"):
                raise HTTPException(422, f"Formula: {e}")
    if body.rate_source == "buildup" and not db.scalar(select(RateBuildup.id).where(RateBuildup.code == body.rate_code)):
        raise HTTPException(422, f"No build-up {body.rate_code}")
    if body.rate_source == "resource" and not db.scalar(select(Resource.id).where(Resource.code == body.rate_code)):
        raise HTTPException(422, f"No resource {body.rate_code}")


@router.post("/jobs/{job_id}/workbook/lines")
def add_line(job_id: int, body: LineIn, db: Session = Depends(get_db), user: User = Depends(current_user)):
    j = _job(db, job_id, user, edit=True)
    _validate_line(db, body)
    data = body.model_dump()
    if data["position"] is None:
        data["position"] = (db.scalar(select(func.max(EstimateLine.position)).where(EstimateLine.job_id == j.id, EstimateLine.section == body.section)) or 0) + 1
    ln = EstimateLine(job_id=j.id, **data)
    db.add(ln)
    db.flush()
    record(db, user.id, "job", j.id, "workbook_add", after={"description": ln.description, "section": ln.section})
    db.commit()
    return {"id": ln.id}


@router.put("/workbook/lines/{wline_id}")
def edit_line(wline_id: int, body: LineIn, db: Session = Depends(get_db), user: User = Depends(current_user)):
    ln = get_or_404(db, EstimateLine, wline_id)
    _job(db, ln.job_id, user, edit=True)
    _validate_line(db, body)
    for k, v in body.model_dump().items():
        if k == "position" and v is None:
            continue
        setattr(ln, k, v)
    db.commit()
    return {"id": ln.id}


@router.delete("/workbook/lines/{wline_id}", status_code=204)
def delete_line(wline_id: int, db: Session = Depends(get_db), user: User = Depends(current_user)):
    ln = get_or_404(db, EstimateLine, wline_id)
    j = _job(db, ln.job_id, user, edit=True)
    if ln.package_id:
        p = db.get(TradePackage, ln.package_id)
        if p:
            p.status, p.awarded_quote_id = "open", None
    record(db, user.id, "job", j.id, "workbook_delete", before={"description": ln.description, "section": ln.section})
    db.delete(ln)
    db.commit()


class VarIn(BaseModel):
    name: str
    value: float
    note: str = ""


@router.put("/jobs/{job_id}/workbook/variables")
def save_variables(job_id: int, body: list[VarIn], db: Session = Depends(get_db), user: User = Depends(current_user)):
    j = _job(db, job_id, user, edit=True)
    names = []
    for v in body:
        try:
            names.append(est.check_name(v.name))
        except ValueError as e:
            raise HTTPException(422, str(e))
    if len(set(names)) != len(names):
        raise HTTPException(422, "Each variable needs its own name")
    existing = {v.name: v for v in db.scalars(select(EstimateVariable).where(EstimateVariable.job_id == j.id))}
    for name, v in zip(names, body):
        row = existing.pop(name, None) or EstimateVariable(job_id=j.id, name=name)
        row.value, row.note = v.value, v.note
        db.add(row)
    for gone in existing.values():
        db.delete(gone)
    db.commit()
    return est.workbook(db, j)["variables"]


class CopyIn(BaseModel):
    from_job_id: int
    sections: list[str] = []


@router.post("/jobs/{job_id}/workbook/copy")
def copy_from(job_id: int, body: CopyIn, db: Session = Depends(get_db), user: User = Depends(current_user)):
    j = _job(db, job_id, user, edit=True)
    src = _job(db, body.from_job_id, user)
    n = est.copy_lines(db, src, j, body.sections or None)
    db.commit()
    return {"copied": n}


@router.get("/workbook/sources")
def copy_sources(db: Session = Depends(get_db), user: User = Depends(current_user)):
    """Jobs with workbook rows, to copy from (templates)."""
    ids = dict(db.execute(select(EstimateLine.job_id, func.count()).group_by(EstimateLine.job_id)).all())
    out = []
    for j in db.scalars(select(Job).where(Job.id.in_(ids))):
        if can_see_job(user, j):
            secs = sorted(set(db.scalars(select(EstimateLine.section).where(EstimateLine.job_id == j.id))))
            out.append({"id": j.id, "name": j.name, "quote_no": j.quote_no, "rows": ids[j.id], "sections": secs})
    return out


# --- packages ---
class PackageIn(BaseModel):
    name: str = Field(min_length=1, max_length=200)
    kind: str = Field(default="subcontract", pattern="^(subcontract|supply)$")
    scope: str = ""
    due: date | None = None
    section: str = "Subcontract"
    stage_id: int | None = None
    margin_pct: float | None = None


def _package(db: Session, pkg_id: int, user: User, edit: bool = False) -> tuple[TradePackage, Job]:
    p = get_or_404(db, TradePackage, pkg_id)
    return p, _job(db, p.job_id, user, edit)


@router.get("/jobs/{job_id}/packages")
def list_packages(job_id: int, db: Session = Depends(get_db), user: User = Depends(current_user)):
    j = _job(db, job_id, user)
    return [est.package_out(db, p) for p in db.scalars(select(TradePackage).where(TradePackage.job_id == j.id).order_by(TradePackage.id))]


@router.post("/jobs/{job_id}/packages")
def add_package(job_id: int, body: PackageIn, db: Session = Depends(get_db), user: User = Depends(current_user)):
    j = _job(db, job_id, user, edit=True)
    p = TradePackage(job_id=j.id, **body.model_dump())
    db.add(p)
    db.commit()
    return est.package_out(db, p)


@router.put("/packages/{pkg_id}")
def edit_package(pkg_id: int, body: PackageIn, db: Session = Depends(get_db), user: User = Depends(current_user)):
    p, _j = _package(db, pkg_id, user, True)
    for k, v in body.model_dump().items():
        setattr(p, k, v)
    if p.awarded_quote_id:
        q = db.get(PackageQuote, p.awarded_quote_id)
        if q:
            est.award(db, p, q)  # keep the workbook row in step
    db.commit()
    return est.package_out(db, p)


@router.delete("/packages/{pkg_id}", status_code=204)
def delete_package(pkg_id: int, db: Session = Depends(get_db), user: User = Depends(current_user)):
    p, _j = _package(db, pkg_id, user, True)
    for ln in db.scalars(select(EstimateLine).where(EstimateLine.package_id == p.id)):
        db.delete(ln)
    for q in db.scalars(select(PackageQuote).where(PackageQuote.package_id == p.id)):
        if q.file_key:
            storage.delete(q.file_key)
    db.delete(p)
    db.commit()


class QuoteIn(BaseModel):
    company: str = Field(min_length=1, max_length=200)
    email: str = ""
    supplier_id: int | None = None
    status: str = Field(default="invited", pattern="^(invited|received|declined)$")
    amount: float | None = Field(default=None, ge=0)
    adjustments: list[dict] = []
    exclusions: str = ""
    valid_until: date | None = None
    notes: str = ""


@router.post("/packages/{pkg_id}/quotes")
def add_quote(pkg_id: int, body: QuoteIn, db: Session = Depends(get_db), user: User = Depends(current_user)):
    p, _j = _package(db, pkg_id, user, True)
    data = body.model_dump()
    if data["amount"] is not None and data["status"] == "invited":
        data["status"] = "received"
    q = PackageQuote(package_id=p.id, **data)
    if q.status == "received":
        q.received_at = datetime.now(timezone.utc)
    db.add(q)
    db.commit()
    return est.package_out(db, p)


def _quote(db: Session, quote_id: int, user: User) -> tuple[PackageQuote, TradePackage]:
    q = get_or_404(db, PackageQuote, quote_id)
    p, _j = _package(db, q.package_id, user, True)
    return q, p


@router.put("/package-quotes/{quote_id}")
def edit_quote(quote_id: int, body: QuoteIn, db: Session = Depends(get_db), user: User = Depends(current_user)):
    q, p = _quote(db, quote_id, user)
    data = body.model_dump()
    adj = []
    for a in data.pop("adjustments") or []:
        try:
            adj.append({"description": str(a.get("description") or "")[:200], "amount": float(a.get("amount") or 0)})
        except (TypeError, ValueError):
            raise HTTPException(422, "Each levelling line needs an amount")
    q.adjustments = adj
    for k, v in data.items():
        setattr(q, k, v)
    if q.amount is not None and q.status == "invited":
        q.status = "received"
    if q.status == "received" and q.received_at is None:
        q.received_at = datetime.now(timezone.utc)
    if p.awarded_quote_id == q.id:
        est.award(db, p, q)
    db.commit()
    return est.package_out(db, p)


@router.delete("/package-quotes/{quote_id}", status_code=204)
def delete_quote(quote_id: int, db: Session = Depends(get_db), user: User = Depends(current_user)):
    q, p = _quote(db, quote_id, user)
    if p.awarded_quote_id == q.id:
        raise HTTPException(409, "That's the awarded quote. Award another first.")
    if q.file_key:
        storage.delete(q.file_key)
    db.delete(q)
    db.commit()


@router.post("/package-quotes/{quote_id}/file")
async def quote_file(quote_id: int, file: UploadFile = File(...), db: Session = Depends(get_db), user: User = Depends(current_user)):
    q, p = _quote(db, quote_id, user)
    data = await file.read(config.MAX_UPLOAD_BYTES + 1)
    if len(data) > config.MAX_UPLOAD_BYTES:
        raise HTTPException(413, "File too large")
    if q.file_key:
        storage.delete(q.file_key)
    name = file.filename or "quote.pdf"
    q.file_key = f"jobs/{p.job_id}/quotes/{uuid.uuid4().hex}-{''.join(c if c.isalnum() or c in '.-_' else '_' for c in name)[:120]}"
    q.file_name = name
    storage.save_bytes(q.file_key, data)
    db.commit()
    return {"ok": True}


@router.get("/package-quotes/{quote_id}/file")
def get_quote_file(quote_id: int, db: Session = Depends(get_db), user: User = Depends(current_user)):
    q = get_or_404(db, PackageQuote, quote_id)
    _package(db, q.package_id, user)
    if not q.file_key:
        raise HTTPException(404, "No file")
    return Response(storage.read_bytes(q.file_key), media_type=mimetypes.guess_type(q.file_name)[0] or "application/octet-stream",
                    headers={"Content-Disposition": f'inline; filename="{"".join(c if c.isalnum() or c in ".-_ " else "_" for c in q.file_name)}"'})


@router.post("/packages/{pkg_id}/invite")
def invite(pkg_id: int, db: Session = Depends(get_db), user: User = Depends(current_user)):
    p, j = _package(db, pkg_id, user, True)
    n = est.invite(db, p, j, user)
    db.commit()
    return {"sent": n}


class AwardIn(BaseModel):
    quote_id: int


@router.post("/packages/{pkg_id}/award")
def award(pkg_id: int, body: AwardIn, db: Session = Depends(get_db), user: User = Depends(current_user)):
    p, j = _package(db, pkg_id, user, True)
    q = get_or_404(db, PackageQuote, body.quote_id)
    if q.package_id != p.id:
        raise HTTPException(404, "Quote not found")
    try:
        ln = est.award(db, p, q)
    except ValueError as e:
        raise HTTPException(409, str(e))
    record(db, user.id, "job", j.id, "package_award", after={"package": p.name, "company": q.company, "amount": est.levelled(q)})
    db.commit()
    return {"line_id": ln.id, **est.package_out(db, p)}


# --- versions ---
class SnapIn(BaseModel):
    label: str = Field(min_length=1, max_length=100)


@router.get("/jobs/{job_id}/versions")
def list_versions(job_id: int, db: Session = Depends(get_db), user: User = Depends(current_user)):
    j = _job(db, job_id, user)
    names = {u.id: u.name or u.email for u in db.scalars(select(User))}
    return [{"id": s.id, "label": s.label, "sell": (s.summary or {}).get("sell_ex_gst"), "cost": (s.summary or {}).get("cost"),
             "by": names.get(s.created_by or 0, ""), "at": s.created_at.isoformat()}
            for s in db.scalars(select(EstimateSnapshot).where(EstimateSnapshot.job_id == j.id).order_by(EstimateSnapshot.created_at.desc()))]


@router.post("/jobs/{job_id}/versions")
def save_version(job_id: int, body: SnapIn, db: Session = Depends(get_db), user: User = Depends(current_user)):
    j = _job(db, job_id, user)
    s = est.snapshot(db, j, body.label, user)
    db.commit()
    return {"id": s.id}


@router.get("/jobs/{job_id}/versions/compare")
def compare(job_id: int, a: str, b: str = "current", db: Session = Depends(get_db), user: User = Depends(current_user)):
    j = _job(db, job_id, user)

    def load(x: str) -> dict:
        if x == "current":
            return est.current_version(db, j)
        s = db.get(EstimateSnapshot, int(x)) if x.isdigit() else None
        if s is None or s.job_id != j.id:
            raise HTTPException(404, "Version not found")
        return est.version_dict(s)
    return est.compare(load(a), load(b))


@router.delete("/estimate-versions/{snap_id}", status_code=204)
def delete_version(snap_id: int, db: Session = Depends(get_db), user: User = Depends(current_user)):
    s = get_or_404(db, EstimateSnapshot, snap_id)
    _job(db, s.job_id, user)
    db.delete(s)
    db.commit()


# --- programme and cash flow ---
class ProgrammeIn(BaseModel):
    start: date
    retention_pct: float = Field(default=5, ge=0, le=20)
    retention_cap_pct: float = Field(default=5, ge=0, le=20)
    terms_days: int = Field(default=30, ge=0, le=180)
    pay_suppliers_days: int = Field(default=30, ge=0, le=180)
    activities: list[dict] = []


@router.get("/jobs/{job_id}/programme")
def get_programme(job_id: int, db: Session = Depends(get_db), user: User = Depends(current_user)):
    j = _job(db, job_id, user)
    p = db.scalar(select(TenderProgramme).where(TenderProgramme.job_id == j.id))
    prog = est.programme_dict(p) if p else est.default_programme(db, j)
    return {"programme": prog, "cashflow": est.cashflow(db, j, prog)}


@router.put("/jobs/{job_id}/programme")
def save_programme(job_id: int, body: ProgrammeIn, db: Session = Depends(get_db), user: User = Depends(current_user)):
    j = _job(db, job_id, user)
    acts = []
    for a in body.activities:
        curve = a.get("curve") or "s"
        if curve not in est.CURVES:
            raise HTTPException(422, f"Curve must be one of {', '.join(est.CURVES)}")
        acts.append({"name": str(a.get("name") or "Activity")[:100], "stages": [str(s) for s in a.get("stages") or []],
                     "start_week": max(0, int(a.get("start_week") or 0)), "weeks": max(1, min(260, int(a.get("weeks") or 1))), "curve": curve})
    p = db.scalar(select(TenderProgramme).where(TenderProgramme.job_id == j.id)) or TenderProgramme(job_id=j.id, start=body.start)
    p.start, p.retention_pct, p.retention_cap_pct = body.start, body.retention_pct, body.retention_cap_pct
    p.terms_days, p.pay_suppliers_days, p.activities = body.terms_days, body.pay_suppliers_days, acts
    db.add(p)
    db.commit()
    prog = est.programme_dict(p)
    return {"programme": prog, "cashflow": est.cashflow(db, j, prog)}


@router.post("/jobs/{job_id}/programme/to-board")
def programme_to_board(job_id: int, db: Session = Depends(get_db), user: User = Depends(current_user)):
    """Turn the programme into a project programme board (timeline view), e.g. once the job is won."""
    from .. import boards as bd

    j = _job(db, job_id, user)
    p = db.scalar(select(TenderProgramme).where(TenderProgramme.job_id == j.id))
    prog = est.programme_dict(p) if p else est.default_programme(db, j)
    cf = est.cashflow(db, j, prog)
    b = bd.from_template(db, "programme", user, f"{j.name} programme", job_id=j.id)
    for g in cf["gantt"]:
        bd.create_item(db, b, g["name"], user.id, values={"when": {"start": g["start"], "end": g["finish"]}, "status": "todo", "job": j.id})
    db.commit()
    return {"board_id": b.id}


# accept a form post with a file for the quote in one go (from the phone or a quick add)
@router.post("/packages/{pkg_id}/quotes/upload")
async def add_quote_with_file(pkg_id: int, company: str = Form(...), amount: float | None = Form(None), email: str = Form(""),
                              exclusions: str = Form(""), file: UploadFile | None = File(None), db: Session = Depends(get_db),
                              user: User = Depends(current_user)):
    p, _j = _package(db, pkg_id, user, True)
    q = PackageQuote(package_id=p.id, company=company, amount=amount, email=email, exclusions=exclusions,
                     status="received" if amount is not None else "invited",
                     received_at=datetime.now(timezone.utc) if amount is not None else None)
    db.add(q)
    db.flush()
    if file is not None and file.filename:
        data = await file.read(config.MAX_UPLOAD_BYTES + 1)
        q.file_key = f"jobs/{p.job_id}/quotes/{uuid.uuid4().hex}-{''.join(c if c.isalnum() or c in '.-_' else '_' for c in file.filename)[:120]}"
        q.file_name = file.filename
        storage.save_bytes(q.file_key, data)
    db.commit()
    return est.package_out(db, p)
